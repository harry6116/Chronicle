import os, sys, glob, json, time, re, threading, platform, subprocess, logging, shutil, webbrowser, random, html, csv, io, hashlib, traceback, faulthandler
from collections import OrderedDict
import base64

try:
    import cv2
except ImportError:  # pragma: no cover - optional dependency in lightweight test envs
    cv2 = None

try:
    from PIL import Image
except ImportError:  # pragma: no cover - optional dependency in lightweight test envs
    Image = None

try:
    import fitz
except ImportError:  # pragma: no cover - optional dependency in lightweight test envs
    fitz = None

try:
    import docx
except ImportError:  # pragma: no cover - optional dependency in lightweight test envs
    docx = None

try:
    from fpdf import FPDF
except ImportError:  # pragma: no cover - optional dependency in lightweight test envs
    FPDF = None

try:
    from google import genai
except ImportError:  # pragma: no cover - optional dependency in lightweight test envs
    genai = None

try:
    import openpyxl
except ImportError:  # pragma: no cover - optional dependency in lightweight test envs
    openpyxl = None

try:
    from ebooklib import epub
except ImportError:  # pragma: no cover - optional dependency in lightweight test envs
    epub = None

import wx
import wx.adv as wxadv
import signal
try:
    import winsound
except ImportError:
    winsound = None
from chronicle_core import (
    clean_text_artifacts as core_clean_text_artifacts,
    apply_modern_punctuation as core_apply_modern_punctuation,
    apply_modern_currency as core_apply_modern_currency,
    apply_expanded_abbreviations as core_apply_expanded_abbreviations,
    apply_output_integrity_contract as core_apply_output_integrity_contract,
    csv_to_accessible_text as core_csv_to_accessible_text,
    sanitize_latin1 as core_sanitize_latin1,
    sanitize_model_output as core_sanitize_model_output,
    normalize_streamed_html_document as core_normalize_streamed_html_document,
    write_header as core_write_header,
    write_footer as core_write_footer,
)
from chronicle_app.services.prompting import (
    build_prompt as build_shared_prompt,
    enforce_archival_heading_structure as shared_enforce_archival_heading_structure,
    get_output_lang_code as shared_get_output_lang_code,
    get_output_text_direction as shared_get_output_text_direction,
    get_translation_target as shared_get_translation_target,
    strip_synthetic_page_filename_headings as shared_strip_synthetic_page_filename_headings,
)
from chronicle_app.config import (
    MODEL_LABEL_BY_KEY,
    PROFILE_CHOICES,
    PROFILE_KEY_TO_LABEL,
    PROFILE_LABEL_TO_KEY,
    PROFILE_PRESETS,
    RTL_LANGUAGE_CODES,
    TRANSLATION_TARGETS,
)
from chronicle_app.services.runtime_policies import (
    DEFAULT_CLAUDE_MODEL,
    build_profile_selection_summary as shared_build_profile_selection_summary,
    get_preferred_profile_model as shared_get_preferred_profile_model,
    get_processing_speed_warning as shared_get_processing_speed_warning,
    get_model_tradeoff_text as shared_get_model_tradeoff_text,
    get_pdf_chunk_pages as shared_get_pdf_chunk_pages,
    persist_runtime_settings_to_cfg as shared_persist_runtime_settings_to_cfg,
    resolve_model_for_available_keys as shared_resolve_model_for_available_keys,
    wait_for_gemini_upload_ready as shared_wait_for_gemini_upload_ready,
)
from chronicle_app.services.app_files import (
    build_log_header as shared_build_log_header,
    emit_launch_continuity as shared_emit_launch_continuity,
    get_runtime_build_stamp as shared_get_runtime_build_stamp,
    load_json_file as shared_load_json_file,
    resolve_runtime_crash_log_path as shared_resolve_runtime_crash_log_path,
    resolve_default_log_directory as shared_resolve_default_log_directory,
    save_json_file as shared_save_json_file,
    update_continuity_runtime_status as shared_update_continuity_runtime_status,
    write_processing_log as shared_write_processing_log,
)
try:
    from chronicle_app.services.licensing import (
        format_license_status as shared_format_license_status,
        install_license_file as shared_install_license_file,
        load_installed_license as shared_load_installed_license,
        resolve_public_key as shared_resolve_public_key,
    )
except ImportError:  # pragma: no cover - optional in bare system python before deps are installed
    shared_format_license_status = None
    shared_install_license_file = None
    shared_load_installed_license = None
    shared_resolve_public_key = None
from chronicle_app.services.document_processors import (
    estimate_text_work_units as shared_estimate_text_work_units,
    process_epub as shared_process_epub,
    process_img as shared_process_img,
    process_pptx as shared_process_pptx,
    process_text as shared_process_text,
)
from chronicle_app.services.pdf_processor import process_pdf as shared_process_pdf
from chronicle_app.services.exporters import (
    dispatch_save as shared_dispatch_save,
    save_docx as shared_save_docx,
    save_epub as shared_save_epub,
    save_pdf as shared_save_pdf,
)
from chronicle_app.services.processing_runtime import (
    HeartbeatMonitor as SharedHeartbeatMonitor,
    RequestRuntime,
    build_payload as shared_build_payload,
    build_request_cache_key as shared_build_request_cache_key,
    generate_retry as shared_generate_retry,
    handle_stream as shared_handle_stream,
    sha256_file as shared_sha256_file,
    sha256_text as shared_sha256_text,
    stream_with_cache as shared_stream_with_cache,
)
from chronicle_app.services.session_runtime import (
    build_session_payload as shared_build_session_payload,
    delete_active_session_file as shared_delete_active_session_file,
    has_incomplete_items as shared_has_incomplete_items,
    restore_session_queue as shared_restore_session_queue,
    save_active_session_file as shared_save_active_session_file,
)
from chronicle_app.services.scan_runtime import (
    collect_scan_files as _collect_scan_files,
    driver_from_scanner_source as _driver_from_scanner_source,
    merge_scan_files_to_single_pdf as _merge_scan_files_to_single_pdf,
)
from chronicle_app.services.queue_runtime import (
    add_path_entries as shared_add_path_entries,
    collect_supported_files_from_folder as shared_collect_supported_files_from_folder,
    find_queue_rows_by_paths as shared_find_queue_rows_by_paths,
)
from chronicle_app.services.queue_state_runtime import (
    apply_settings_to_rows as shared_apply_settings_to_rows,
    build_progress_summary as shared_build_progress_summary,
    estimate_path_work_units as shared_estimate_path_work_units,
    get_run_unit_totals as shared_get_run_unit_totals,
    get_target_queue_indices_for_setting_change as shared_get_target_queue_indices_for_setting_change,
    pause_selected_tasks as shared_pause_selected_tasks,
    refresh_queue_work_unit_estimates as shared_refresh_queue_work_unit_estimates,
    resume_selected_tasks as shared_resume_selected_tasks,
    should_log_page_progress as shared_should_log_page_progress,
    should_status_echo_log as shared_should_status_echo_log,
    stop_selected_tasks as shared_stop_selected_tasks,
)
from chronicle_app.services.run_control_runtime import (
    build_running_state_update as shared_build_running_state_update,
    count_active_queue_items as shared_count_active_queue_items,
    count_saved_queue_items as shared_count_saved_queue_items,
    pause_current_processing_row as shared_pause_current_processing_row,
    prepare_running_close as shared_prepare_running_close,
    wait_while_paused as shared_wait_while_paused,
)
from chronicle_app.services.worker_runtime import (
    MirroredProgressMemory,
    MirroredTextWriter,
    build_progress_state_header,
    build_worker_run_plan as shared_build_worker_run_plan,
    compute_target_dir as shared_compute_target_dir,
    determine_needs_pdf_audit as shared_determine_needs_pdf_audit,
    ensure_progress_sidecar_header,
    estimate_current_file_total_units as shared_estimate_current_file_total_units,
    load_merge_resume_state as shared_load_merge_resume_state,
    prepare_job_execution_context as shared_prepare_job_execution_context,
    resolve_output_path as shared_resolve_output_path,
    write_progress_state,
)
from chronicle_app.services.worker_finalize_runtime import (
    append_pdf_audit_appendix_if_needed as shared_append_pdf_audit_appendix_if_needed,
    cleanup_output_text as shared_cleanup_output_text,
    finalize_job_success as shared_finalize_job_success,
    finalize_merged_output as shared_finalize_merged_output,
    finalize_worker_completion as shared_finalize_worker_completion,
    finalize_worker_session as shared_finalize_worker_session,
    handle_job_error as shared_handle_job_error,
)
from chronicle_app.services.worker_execute_runtime import (
    process_job_content as shared_process_job_content,
)
from chronicle_app.services.adaptive_engine_routing import select_execution_model_for_job
from chronicle_app.ui.bindings import bind_named
from chronicle_app.services.scan_flow_runtime import (
    apply_scan_settings as shared_apply_scan_settings,
    begin_scan_session as shared_begin_scan_session,
    build_scan_completion_message as shared_build_scan_completion_message,
    build_scan_start_message as shared_build_scan_start_message,
    choose_scan_commands as shared_choose_scan_commands,
    execute_scan_commands as shared_execute_scan_commands,
    resolve_scan_driver as shared_resolve_scan_driver,
    resolve_scan_output_files as shared_resolve_scan_output_files,
)
from chronicle_app.services.ordering_runtime import (
    get_ordered_jobs_for_processing as shared_get_ordered_jobs_for_processing,
    get_page_sequence_number as shared_get_page_sequence_number,
    resolve_merge_output_path as shared_resolve_merge_output_path,
)
from chronicle_app.services.scheduling_runtime import (
    build_schedule_summary_label as shared_build_schedule_summary_label,
    format_timestamp_local as shared_format_timestamp_local,
    normalize_future_timestamp as shared_normalize_future_timestamp,
    should_trigger_scheduled_start as shared_should_trigger_scheduled_start,
)
from chronicle_app.services.run_start_runtime import (
    apply_start_configuration as shared_apply_start_configuration,
    begin_run_start as shared_begin_run_start,
    build_run_reset_state as shared_build_run_reset_state,
    build_start_messages as shared_build_start_messages,
    collect_pending_rows as shared_collect_pending_rows,
    expand_multi_range_pdf_rows as shared_expand_multi_range_pdf_rows,
    find_missing_api_key_requirement as shared_find_missing_api_key_requirement,
    prepare_queue_for_start as shared_prepare_queue_for_start,
    validate_output_destination as shared_validate_output_destination,
    validate_pending_pdf_page_scopes as shared_validate_pending_pdf_page_scopes,
)
from chronicle_app.services.scanner_discovery import (
    candidate_naps2_commands as _candidate_naps2_commands,
    discover_connected_flatbed_scanners,
    discover_scanners_naps2,
    run_command_capture as _run_command_capture,
)
try:
    from chronicle_app.ui.dialogs import (
        ApiKeyDialog,
        CloseRunningDialog,
        FirstLaunchNoticeDialog,
        PrefsDialog,
        ScanSettingsDialog,
        ScheduleExtractionDialog,
        SessionRecoveryDialog,
    )
    from chronicle_app.ui.menus import ChronicleMenuBar
    from chronicle_app.ui.queue_panel import QueuePanel
except Exception:  # pragma: no cover - allows lightweight unit tests with wx stubs
    ApiKeyDialog = None
    CloseRunningDialog = None
    FirstLaunchNoticeDialog = None
    PrefsDialog = None
    ScanSettingsDialog = None
    ScheduleExtractionDialog = None
    SessionRecoveryDialog = None
    ChronicleMenuBar = None
    QueuePanel = None
from chronicle_app.ui.queue_support import (
    build_queue_accessibility_description,
    build_queue_accessibility_name,
    build_queue_current_row_announcement,
    ensure_queue_table_landing,
    get_queue_display_status,
)
from chronicle_app.ui.main_frame_sections import (
    build_log_section,
    build_progress_section,
    build_queue_action_section,
    build_settings_section,
)


class _FitzPageAdapter:
    def __init__(self, document, index):
        self._document = document
        self._index = index

    def extract_text(self):
        return self._document.load_page(self._index).get_text("text")


class PdfReader:
    def __init__(self, path):
        if fitz is None:
            raise RuntimeError("PyMuPDF (fitz) is required for PDF reading.")
        self._path = str(path)
        self._document = fitz.open(self._path)
        self.pages = [_FitzPageAdapter(self._document, idx) for idx in range(len(self._document))]


class PdfWriter:
    def __init__(self):
        self._source_pages = []

    def add_page(self, page):
        self._source_pages.append((page._document, page._index))

    def write(self, fh):
        if fitz is None:
            raise RuntimeError("PyMuPDF (fitz) is required for PDF writing.")
        output_doc = fitz.open()
        try:
            for source_doc, page_index in self._source_pages:
                output_doc.insert_pdf(source_doc, from_page=page_index, to_page=page_index)
            output_doc.save(fh, garbage=3, deflate=True)
        finally:
            output_doc.close()


# --- PYINSTALLER & MAC ENVIRONMENT HACKS ---
brew_paths = glob.glob("/opt/homebrew/lib/python3.*/site-packages") + glob.glob("/usr/local/lib/python3.*/site-packages")
for path in brew_paths:
    if os.path.exists(path) and path not in sys.path: sys.path.append(path)

logging.getLogger("fitz").setLevel(logging.ERROR)

_CRASH_LOG_STREAM = None


def _install_runtime_crash_capture():
    global _CRASH_LOG_STREAM
    try:
        crash_log_path = shared_resolve_runtime_crash_log_path()
        _CRASH_LOG_STREAM = open(crash_log_path, 'a', buffering=1, encoding='utf-8', errors='replace')
        print(f"\n[{time.strftime('%Y-%m-%d %H:%M:%S')}] Chronicle crash capture initialized.\n", file=_CRASH_LOG_STREAM)
        try:
            faulthandler.enable(file=_CRASH_LOG_STREAM, all_threads=True)
        except Exception:
            pass

        def _write_uncaught(prefix, exc_type, exc_value, exc_traceback):
            try:
                print(
                    f"[{time.strftime('%Y-%m-%d %H:%M:%S')}] {prefix}: {exc_type.__name__}: {exc_value}",
                    file=_CRASH_LOG_STREAM,
                )
                traceback.print_exception(exc_type, exc_value, exc_traceback, file=_CRASH_LOG_STREAM)
                _CRASH_LOG_STREAM.flush()
            except Exception:
                pass

        def _sys_excepthook(exc_type, exc_value, exc_traceback):
            _write_uncaught("Unhandled exception", exc_type, exc_value, exc_traceback)
            sys.__excepthook__(exc_type, exc_value, exc_traceback)

        sys.excepthook = _sys_excepthook

        if hasattr(threading, "excepthook"):
            def _thread_excepthook(args):
                thread_name = getattr(getattr(args, "thread", None), "name", "thread")
                _write_uncaught(f"Unhandled thread exception in {thread_name}", args.exc_type, args.exc_value, args.exc_traceback)
            threading.excepthook = _thread_excepthook
    except Exception:
        pass


_install_runtime_crash_capture()

if getattr(sys, 'frozen', False):
    try:
        crash_log_path = shared_resolve_runtime_crash_log_path()
        sys.stdout = _CRASH_LOG_STREAM or open(crash_log_path, 'a', buffering=1, encoding='utf-8', errors='replace')
        sys.stderr = sys.stdout
    except Exception:
        pass

# --- CONSTANTS ---
APP_NAME = "Chronicle"
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))


def _get_gui_app_data_dir():
    if platform.system() == "Windows":
        base = os.environ.get("APPDATA") or os.path.expanduser("~")
        return os.path.join(base, APP_NAME)
    if platform.system() == "Darwin":
        return os.path.join(os.path.expanduser("~/Library/Application Support"), APP_NAME)
    base = os.environ.get("XDG_CONFIG_HOME") or os.path.expanduser("~/.config")
    return os.path.join(base, APP_NAME)


APP_DATA_DIR = _get_gui_app_data_dir()
os.makedirs(APP_DATA_DIR, exist_ok=True)
KEYS_FILE = os.path.join(APP_DATA_DIR, "api_keys.json")
CONFIG_FILE = os.path.join(APP_DATA_DIR, "user_config.json")
SESSION_FILE = os.path.join(APP_DATA_DIR, "chronicle_active_session.json")
LICENSE_FILE = os.path.join(APP_DATA_DIR, "chronicle-license.json")
HELP_FILE = os.path.join(SCRIPT_DIR, "docs", "user", "chronicle_help.html")
DONATE_BUYMEACOFFEE_URL = "https://buymeacoffee.com/thevoiceguy"
DONATE_PAYPAL_URL = "https://paypal.me/MarshallVoiceovers"
AUTORUN_SPEC_ENV = "CHRONICLE_AUTORUN_SPEC"
AUTORUN_CONTINUITY_DIR_ENV = "CHRONICLE_CONTINUITY_DIR"
SESSION_SCHEMA_VERSION = 1
PDF_CHUNK_PAGES = 3
TEXT_CHUNK_CHARS = 10000
API_MIN_REQUEST_INTERVAL_SEC = 1.0
API_BACKOFF_BASE_SEC = 8
API_BACKOFF_MAX_SEC = 90
API_MAX_RETRIES = 6
API_MAX_CONCURRENT_REQUESTS = 1
API_MAX_PENDING_REQUESTS = 24
API_REQUEST_QUEUE_POLL_SEC = 0.25
API_CACHE_MAX_ENTRIES = 3000
TEXT_BATCH_TARGET_CHARS = 18000
PDF_TEXTLAYER_AUDIT_COVERAGE_WARN = 0.72
PDF_TEXTLAYER_AUDIT_COVERAGE_APPEND_FULL = 0.60
PDF_TEXTLAYER_AUDIT_MAX_SOURCE_CHARS = 1_500_000
PDF_TEXTLAYER_AUDIT_MAX_OUTPUT_CHARS = 1_500_000
PDF_TEXTLAYER_AUDIT_MAX_MISSING_LINES = 800
LOW_MEMORY_PDF_AUDIT_SKIP_MB = 40
SESSION_TERMINAL_STATUSES = {"Done", "Error", "Skipped", "Unsupported", "Missing", "Stopped"}
QUEUE_EMPTY_PLACEHOLDER = "Queue is empty. Use Add Files or Add Folder to load items."
STREAMABLE_FORMATS = {"html", "txt", "md"}
SUPPORTED_EXTENSIONS = ['.pdf', '.docx', '.txt', '.md', '.rtf', '.csv', '.js', '.jpg', '.jpeg', '.png', '.bmp', '.tiff', '.tif', '.webp', '.xlsx', '.xls', '.pptx', '.ppt', '.epub']
PROTECTED_INPUT_DIR_BASENAMES = {"checking documents"}
ROW_SETTING_KEYS = (
    "format_type",
    "doc_profile",
    "model_name",
    "model_override",
    "translate_mode",
    "translate_target",
    "modernize_punctuation",
    "unit_conversion",
    "abbrev_expansion",
    "image_descriptions",
    "preserve_original_page_numbers",
    "merge_files",
    "large_print",
    "pdf_page_scope",
)
_api_request_lock = threading.Lock()
_last_api_request_ts = 0.0
_api_request_semaphore = threading.Semaphore(API_MAX_CONCURRENT_REQUESTS)
_api_queue_lock = threading.Lock()
_api_pending_requests = 0
_chunk_cache = OrderedDict()
_chunk_cache_lock = threading.Lock()
ACTIVE_FRAME = None
request_runtime = RequestRuntime(
    api_min_request_interval_sec=API_MIN_REQUEST_INTERVAL_SEC,
    api_max_pending_requests=API_MAX_PENDING_REQUESTS,
    api_request_queue_poll_sec=API_REQUEST_QUEUE_POLL_SEC,
    api_max_concurrent_requests=API_MAX_CONCURRENT_REQUESTS,
    api_cache_max_entries=API_CACHE_MAX_ENTRIES,
)

class StopRequestedError(Exception):
    pass

def bool_text(value, true_text, false_text):
    return true_text if value else false_text

def pace_api_request(log_cb=print):
    return request_runtime.pace_api_request(log_cb=log_cb)

def sha256_text(text):
    return shared_sha256_text(text)

def sha256_file(path):
    return shared_sha256_file(path)

def build_request_cache_key(model, prompt, payload_kind, payload_fingerprint):
    return shared_build_request_cache_key(model, prompt, payload_kind, payload_fingerprint)

def cache_get(cache_key):
    return request_runtime.cache_get(cache_key)

def cache_put(cache_key, text):
    return request_runtime.cache_put(cache_key, text)

def wait_for_request_slot(log_cb=print):
    return request_runtime.wait_for_request_slot(log_cb=log_cb)

def release_request_slot():
    return request_runtime.release_request_slot()

def batch_text_chunks(chunks, target_chars=TEXT_BATCH_TARGET_CHARS):
    batches = []
    current_batch = []
    current_len = 0
    for chunk in chunks:
        if not chunk:
            continue
        chunk_len = len(chunk)
        if current_batch and (current_len + chunk_len) > target_chars:
            batches.append("\n".join(current_batch))
            current_batch = [chunk]
            current_len = chunk_len
        else:
            current_batch.append(chunk)
            current_len += chunk_len
    if current_batch:
        batches.append("\n".join(current_batch))
    return batches

def legacy_punctuation_text(enabled):
    if enabled:
        return "Normalize legacy punctuation (old-style)"
    return "Keep legacy punctuation (old-style) as-is"

def profile_tooltip_text(profile_key):
    guides = {
        "standard": "Use this when the batch is mixed, miscellaneous, or genuinely hard to classify. This is the fallback preset when no other label clearly fits.",
        "letters": "Use this for personal or business correspondence: letters, memos, notices, circulars, and sign-off-driven documents. Do not use it for full reports.",
        "office": "Use this for ordinary business files such as reports, meeting packs, proposals, briefings, and Word exports. Do not use it for letters or formal government publications.",
        "government": "Use this for public-sector reports, consultation papers, committee papers, appendices, and official record PDFs with repeated government headers or footers.",
        "legal": "Use this for contracts, legislation, regulations, policies, and other clause-heavy legal documents.",
        "forms": "Use this for forms, checklists, worksheets, boxes, signatures, and fill-in pages where blanks matter.",
        "tabular": "Use this for tables, spreadsheets, CSV exports, registers, and row-and-column documents.",
        "manual": "Use this for manuals, procedures, instructions, SOPs, maintenance guides, and technical how-to documents.",
        "slides": "Use this for slide decks, presentations, lecture slides, and speaker handouts.",
        "flyer": "Use this for flyers, posters, one-page notices, event sheets, and short announcements.",
        "brochure": "Use this for brochures, catalogues, pamphlets, and folded or multi-panel handouts.",
        "book": "Use this for books, novels, memoirs, and other long-form prose.",
        "newspaper": "Use this for newspapers and other dense multi-column pages.",
        "academic": "Use this for research papers, journal articles, citations, equations, and footnotes.",
        "transcript": "Use this for interviews, hearings, scripts, speaker turns, and dialogue-heavy text.",
        "poetry": "Use this when line breaks, indentation, and stanza shape must be preserved.",
        "handwritten": "Use this for handwritten letters, personal notes, diaries, drafts, and photographed handwritten pages. Use this instead of Medical unless the page is clearly clinical.",
        "archival": "Use this for historical papers, archives, manuscripts, ledgers, and older source material.",
        "medical": "Use this for clinical notes, referral letters, charts, forms, and doctor handwriting when the content is clearly medical.",
        "military": "Use this for war diaries, operational orders, service records, and other military files.",
        "intelligence": "Use this for cables, signals traffic, intelligence briefings, routing headers, and codewords.",
        "museum": "Use this for object labels, captions, wall text, provenance notes, and exhibit metadata.",
    }
    preset = PROFILE_PRESETS.get(profile_key, PROFILE_PRESETS["standard"])
    label = PROFILE_KEY_TO_LABEL.get(profile_key, "Standard")
    model_label = MODEL_LABEL_BY_KEY.get(preset.get("model_name", "gemini-2.5-flash"), "Gemini 2.5 Flash")
    translation_label = {
        "none": "Disabled",
        "both": "Keep original + bracketed translation",
        "english_only": "Translate only",
    }.get(preset.get("translate_mode", "none"), "Disabled")
    return (
        f"{label}: {guides.get(profile_key, guides['standard'])}\n"
        f"Recommended model (optional): {model_label}\n"
        f"Translation: {translation_label}\n"
        f"Legacy punctuation: {legacy_punctuation_text(preset.get('modernize_punctuation', False))}\n"
        f"Units: {bool_text(preset.get('unit_conversion', False), 'Convert in brackets', 'Keep original')}\n"
        f"Abbreviations: {bool_text(preset.get('abbrev_expansion', False), 'Expand in brackets', 'Keep original')}\n"
        f"Image descriptions: {bool_text(preset.get('image_descriptions', True), 'Enabled', 'Disabled')}"
    )


def get_model_tradeoff_text(model_name):
    return shared_get_model_tradeoff_text(model_name)


def get_processing_speed_warning(profile_key, model_name):
    return shared_get_processing_speed_warning(profile_key, model_name)


def build_profile_selection_summary(profile_key, current_model_name):
    return shared_build_profile_selection_summary(
        profile_key,
        current_model_name,
        profile_label_map=PROFILE_KEY_TO_LABEL,
        profile_presets=PROFILE_PRESETS,
    )


def persist_runtime_settings_to_cfg(cfg, settings):
    return shared_persist_runtime_settings_to_cfg(cfg, settings)


def get_preferred_profile_model(profile_key, cfg):
    return shared_get_preferred_profile_model(
        profile_key,
        cfg=cfg,
        profile_presets=PROFILE_PRESETS,
    )


def get_pdf_chunk_pages(model_name, doc_profile, total_pages, *, file_size_mb=None):
    return shared_get_pdf_chunk_pages(model_name, doc_profile, total_pages, file_size_mb=file_size_mb)


def normalize_pdf_page_scope_text(scope_text):
    text = str(scope_text or "").strip()
    if not text:
        return ""
    if text.lower() in {"all", "*"}:
        return text
    text = re.sub(r"[\r\n;]+", ",", text)
    text = re.sub(r"(?<=\d)\s+(?=\d)", ",", text)
    text = re.sub(r"\s*-\s*", "-", text)
    text = re.sub(r"\s*,\s*", ",", text)
    text = re.sub(r",+", ",", text).strip(", ")
    return text


def parse_pdf_page_scope_spec(scope_text, total_pages):
    normalized = normalize_pdf_page_scope_text(scope_text)
    if total_pages <= 0:
        return []
    if not normalized or normalized.lower() in {"all", "*"}:
        return list(range(total_pages))

    selected = []
    seen = set()
    for raw_part in normalized.split(','):
        part = raw_part.strip()
        if not part:
            raise ValueError("Use page numbers like 1,3,5-7. Empty entries are not allowed.")
        if '-' in part:
            start_txt, end_txt = [piece.strip() for piece in part.split('-', 1)]
            if not start_txt or not end_txt:
                raise ValueError(f"Invalid page range '{part}'. Use start-end, for example 2-5.")
            if not start_txt.isdigit() or not end_txt.isdigit():
                raise ValueError(f"Invalid page range '{part}'. Only numbers and commas are allowed.")
            start = int(start_txt)
            end = int(end_txt)
            if start < 1 or end < 1:
                raise ValueError("Page numbers must start at 1.")
            if end < start:
                raise ValueError(f"Invalid page range '{part}'. The end page must be greater than or equal to the start page.")
            if end > total_pages:
                raise ValueError(f"Page range '{part}' exceeds this PDF's length of {total_pages} pages.")
            for page_num in range(start, end + 1):
                index = page_num - 1
                if index not in seen:
                    selected.append(index)
                    seen.add(index)
            continue
        if not part.isdigit():
            raise ValueError(f"Invalid page value '{part}'. Use numbers like 1,3,5-7.")
        page_num = int(part)
        if page_num < 1:
            raise ValueError("Page numbers must start at 1.")
        if page_num > total_pages:
            raise ValueError(f"Page {page_num} exceeds this PDF's length of {total_pages} pages.")
        index = page_num - 1
        if index not in seen:
            selected.append(index)
            seen.add(index)

    if not selected:
        raise ValueError("No PDF pages were selected.")
    return selected


PDF_PAGE_SCOPE_PRESET_LABELS = ["All Pages", "Page 1", "Pages 1-5", "Custom..."]


def is_valid_pdf_page_scope_text(scope_text):
    normalized = normalize_pdf_page_scope_text(scope_text)
    if not normalized:
        return False
    for raw_part in normalized.split(','):
        part = raw_part.strip()
        if not part:
            return False
        if '-' in part:
            start_txt, end_txt = [piece.strip() for piece in part.split('-', 1)]
            if not start_txt or not end_txt or not start_txt.isdigit() or not end_txt.isdigit():
                return False
            start = int(start_txt)
            end = int(end_txt)
            if start < 1 or end < 1 or end < start:
                return False
            continue
        if not part.isdigit() or int(part) < 1:
            return False
    return True


def pdf_page_scope_value_from_choice_label(label):
    label = str(label or "").strip()
    if label == "All Pages":
        return ""
    if label == "Page 1":
        return "1"
    if label == "Pages 1-5":
        return "1-5"
    if label == "Custom...":
        return ""
    return normalize_pdf_page_scope_text(label)


def pdf_page_scope_choice_label_from_value(scope_text):
    normalized = normalize_pdf_page_scope_text(scope_text)
    if not normalized:
        return "All Pages"
    if normalized == "1":
        return "Page 1"
    if normalized == "1-5":
        return "Pages 1-5"
    return normalized


def wait_for_gemini_upload_ready(client, uploaded, poll_sec=2.5, max_wait_sec=180.0, time_fn=time.time, sleep_fn=time.sleep, log_cb=print):
    return shared_wait_for_gemini_upload_ready(
        client,
        uploaded,
        poll_sec=poll_sec,
        max_wait_sec=max_wait_sec,
        time_fn=time_fn,
        sleep_fn=sleep_fn,
        log_cb=log_cb,
    )

def apply_profile_preset(cfg, profile_key, selected_model_name=None, keep_selected_model=True):
    preset = PROFILE_PRESETS.get(profile_key, PROFILE_PRESETS["standard"])
    current_model = selected_model_name or cfg.get("model_name") or preset.get("model_name", "gemini-2.5-flash")
    cfg.update(preset)
    if keep_selected_model:
        cfg["model_name"] = current_model
    cfg["doc_profile"] = profile_key
    cfg["academic_mode"] = profile_key == "academic"
    return cfg


def build_control_settings_for_profile(cfg, profile_key, *, selected_model_name=None):
    preset_cfg = dict(cfg)
    apply_profile_preset(
        preset_cfg,
        profile_key,
        selected_model_name=selected_model_name,
        keep_selected_model=True,
    )
    return {
        "translate_mode": preset_cfg.get("translate_mode", "none"),
        "modernize_punctuation": bool(preset_cfg.get("modernize_punctuation", False)),
        "unit_conversion": bool(preset_cfg.get("unit_conversion", False)),
        "abbrev_expansion": bool(preset_cfg.get("abbrev_expansion", False)),
        "image_descriptions": bool(preset_cfg.get("image_descriptions", True)),
        "preserve_original_page_numbers": bool(preset_cfg.get("preserve_original_page_numbers", False)),
        "large_print": bool(preset_cfg.get("large_print", False)),
        "merge_files": bool(preset_cfg.get("merge_files", False)),
    }

def get_translation_target(cfg):
    return shared_get_translation_target(cfg, TRANSLATION_TARGETS)

def is_path_within_protected_input_dirs(path):
    try:
        target = os.path.normcase(os.path.abspath(path))
    except Exception:
        return False
    for dirname in PROTECTED_INPUT_DIR_BASENAMES:
        root = os.path.normcase(os.path.abspath(os.path.join(SCRIPT_DIR, dirname)))
        try:
            if os.path.commonpath([target, root]) == root:
                return True
        except Exception:
            continue
    return False

def get_output_lang_code(cfg):
    return shared_get_output_lang_code(cfg, TRANSLATION_TARGETS)

def get_output_text_direction(cfg):
    return shared_get_output_text_direction(cfg, TRANSLATION_TARGETS, RTL_LANGUAGE_CODES)

def load_json(filepath, default_val):
    return shared_load_json_file(filepath, default_val)

def save_json(filepath, data):
    return shared_save_json_file(filepath, data)


def migrate_legacy_gui_app_data():
    legacy_paths = {
        os.path.join(SCRIPT_DIR, "api_keys.json"): KEYS_FILE,
        os.path.join(SCRIPT_DIR, "user_config.json"): CONFIG_FILE,
        os.path.join(SCRIPT_DIR, "chronicle_active_session.json"): SESSION_FILE,
        os.path.join(SCRIPT_DIR, "chronicle-license.json"): LICENSE_FILE,
    }

    for legacy_path, target_path in legacy_paths.items():
        if os.path.abspath(legacy_path) == os.path.abspath(target_path):
            continue
        if not os.path.exists(legacy_path) or os.path.exists(target_path):
            continue
        try:
            shutil.copy2(legacy_path, target_path)
        except Exception:
            continue
        try:
            os.remove(legacy_path)
        except Exception:
            pass


migrate_legacy_gui_app_data()

def get_runtime_build_stamp():
    return shared_get_runtime_build_stamp(
        script_dir=SCRIPT_DIR,
        module_file=__file__,
        sys_executable=sys.executable,
        is_frozen=getattr(sys, "frozen", False),
    )

def build_log_header(build_stamp):
    return shared_build_log_header(build_stamp)

# --- CORE LOGIC & SCRUBBERS ---
HeartbeatMonitor = SharedHeartbeatMonitor
heartbeat = HeartbeatMonitor(exit_fn=os._exit)

def enhance_image(file_path):
    try:
        img = cv2.imread(file_path)
        if img is None: return file_path
        width, height = int(img.shape[1] * 2), int(img.shape[0] * 2)
        resized = cv2.resize(img, (width, height), interpolation=cv2.INTER_LANCZOS4)
        gray = cv2.cvtColor(resized, cv2.COLOR_BGR2GRAY)
        temp_path = file_path + "_enhanced.png"
        cv2.imwrite(temp_path, gray)
        return temp_path
    except: return file_path

def clean_text(text):
    return core_clean_text_artifacts(text)

def csv_to_accessible_text(raw_text, max_rows=None):
    return core_csv_to_accessible_text(raw_text, max_rows=max_rows, max_cell_chars=None)

def sanitize_latin1(text):
    return core_sanitize_latin1(text)

def split_text_for_fail_safe(text, max_chars=900):
    raw = (text or "").strip()
    if not raw:
        return []
    # Sentence-aware split first, then hard-wrap long segments.
    sentence_parts = [p.strip() for p in re.split(r'(?<=[.!?])\s+', raw) if p and p.strip()]
    if not sentence_parts:
        sentence_parts = [raw]
    units = []
    for part in sentence_parts:
        if len(part) <= max_chars:
            units.append(part)
            continue
        start = 0
        while start < len(part):
            units.append(part[start:start + max_chars].strip())
            start += max_chars
    return [u for u in units if u]

def normalize_audit_text(text):
    return re.sub(r'[^a-z0-9]+', '', (text or '').lower())

def describe_quality_score(score_out_of_10, method="vision"):
    if score_out_of_10 >= 8.5:
        base = "Clear page quality."
    elif score_out_of_10 >= 7.0:
        base = "Good quality with minor degradation."
    elif score_out_of_10 >= 5.0:
        base = "Moderate degradation; some text may be difficult."
    else:
        base = "Heavy degradation; risk of missed text."
    if method == "dense-recheck":
        return f"{base} Dense recheck was required."
    if method == "text-layer-fallback":
        return f"{base} Vision fallback was required."
    return base

def assess_image_file_quality(path):
    img = cv2.imread(path)
    if img is None:
        return None
    gray = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)
    mean_luma = float(gray.mean())
    contrast = float(gray.std())
    blur = float(cv2.Laplacian(gray, cv2.CV_64F).var())

    hsv = cv2.cvtColor(img, cv2.COLOR_BGR2HSV)
    # Heuristic stain mask: low saturation + brown/yellow hue ranges often seen in aged/water-damaged pages.
    low_sat = hsv[:, :, 1] < 55
    brownish = ((hsv[:, :, 0] >= 8) & (hsv[:, :, 0] <= 28) & (hsv[:, :, 2] > 55))
    stain_ratio = float((low_sat & brownish).sum()) / max(1, gray.size)

    score = 9.5
    notes = []
    if contrast < 34:
        score -= 1.8
        notes.append("possible faded writing")
    if mean_luma > 205:
        score -= 1.1
        notes.append("washed background / low ink separation")
    if blur < 85:
        score -= 1.3
        notes.append("soft/blurred text edges")
    if stain_ratio > 0.09:
        score -= 1.7
        notes.append("possible water damage or staining artifacts")

    score = max(1.0, min(10.0, score))
    if not notes:
        summary = "Clear page quality."
    else:
        summary = "; ".join(notes).capitalize() + "."
    return score, summary

def get_peak_rss_mb():
    try:
        import resource
        raw = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss
        if raw > 10_000_000:
            return float(raw) / (1024.0 * 1024.0)
        return float(raw) / 1024.0
    except Exception:
        return 0.0

def render_audit_appendix(fmt, heading, body):
    if not body.strip():
        return ""
    if fmt == "html":
        return f"\n<section>\n<h2>{html.escape(heading)}</h2>\n<pre>{html.escape(body)}</pre>\n</section>\n"
    if fmt == "md":
        return f"\n\n## {heading}\n\n```\n{body}\n```\n"
    return f"\n\n[{heading}]\n{body}\n"

def append_generated_text(fmt, file_obj, memory, text):
    if not text:
        return
    if fmt in ['html', 'txt', 'md'] and file_obj:
        file_obj.write(text)
        file_obj.flush()
    if memory is not None:
        memory.append(text)

def strip_synthetic_page_filename_headings(content, fmt):
    return shared_strip_synthetic_page_filename_headings(content, fmt)


def enforce_archival_heading_structure(content, fmt, doc_profile):
    return shared_enforce_archival_heading_structure(content, fmt, doc_profile)

def sanitize_model_output(content, fmt, doc_profile=None, preserve_original_page_numbers=False):
    return core_sanitize_model_output(
        content,
        fmt,
        doc_profile,
        preserve_original_page_numbers,
    )

def apply_modern_punctuation(content):
    return core_apply_modern_punctuation(content)

def apply_modern_currency(content):
    return core_apply_modern_currency(content)


def apply_expanded_abbreviations(content):
    return core_apply_expanded_abbreviations(content)

def normalize_streamed_html_document(full_html):
    return core_normalize_streamed_html_document(full_html)

def apply_output_integrity_contract(content, fmt, doc_profile=None):
    return core_apply_output_integrity_contract(content, fmt, doc_profile)

def cleanup_output_text(content, *, fmt, job_cfg):
    return shared_cleanup_output_text(
        content,
        fmt=fmt,
        job_cfg=job_cfg,
        normalize_html_fn=normalize_streamed_html_document,
        modernize_punctuation_fn=apply_modern_punctuation,
        modernize_currency_fn=apply_modern_currency,
        expand_abbreviations_fn=apply_expanded_abbreviations,
        enforce_heading_structure_fn=enforce_archival_heading_structure,
        apply_integrity_contract_fn=apply_output_integrity_contract,
    )


def run_pdf_textlayer_audit(pdf_path, extracted_text, *, page_scope=""):
    try:
        reader = PdfReader(pdf_path)
    except Exception:
        return None
    total_pages = len(reader.pages)
    selected_page_indices = list(range(total_pages))
    normalized_scope = normalize_pdf_page_scope_text(page_scope)
    if normalized_scope:
        try:
            selected_page_indices = parse_pdf_page_scope_spec(normalized_scope, total_pages)
        except Exception:
            selected_page_indices = list(range(total_pages))
    src_lines = []
    for index in selected_page_indices:
        page = reader.pages[index]
        page_txt = page.extract_text() or ""
        if page_txt.strip():
            src_lines.extend([ln.strip() for ln in page_txt.splitlines() if ln.strip()])
    source_text = "\n".join(src_lines).strip()
    if not source_text:
        return None

    source_truncated = False
    output_truncated = False
    if len(source_text) > PDF_TEXTLAYER_AUDIT_MAX_SOURCE_CHARS:
        source_text = source_text[:PDF_TEXTLAYER_AUDIT_MAX_SOURCE_CHARS]
        source_truncated = True
    extracted_for_audit = extracted_text or ""
    if len(extracted_for_audit) > PDF_TEXTLAYER_AUDIT_MAX_OUTPUT_CHARS:
        extracted_for_audit = extracted_for_audit[:PDF_TEXTLAYER_AUDIT_MAX_OUTPUT_CHARS]
        output_truncated = True

    norm_source = normalize_audit_text(source_text)
    norm_output = normalize_audit_text(extracted_for_audit)
    source_space_norm = " ".join(source_text.lower().split())
    output_space_norm = " ".join(extracted_for_audit.lower().split())
    coverage = min(1.0, len(norm_output) / max(1, len(norm_source)))

    missing = []
    for ln in src_lines:
        plain = " ".join(ln.split())
        if len(plain) < 20:
            continue
        if plain.lower() not in output_space_norm:
            missing.append(plain)
            if len(missing) >= PDF_TEXTLAYER_AUDIT_MAX_MISSING_LINES:
                break

    return {
        "coverage": coverage,
        "source_chars": len(norm_source),
        "output_chars": len(norm_output),
        "missing_lines": missing,
        "source_text": source_text,
        "source_truncated": source_truncated,
        "output_truncated": output_truncated,
    }

def build_prompt(cfg):
    return build_shared_prompt(
        cfg,
        translation_targets=TRANSLATION_TARGETS,
        rtl_language_codes=RTL_LANGUAGE_CODES,
        format_type=cfg.get("format_type", "html"),
    )

# --- FILE OUTPUT EXPORTERS ---
def write_header(file_obj, title, fmt, lang_code="en", text_dir="ltr"):
    core_write_header(file_obj, title, fmt, lang_code=lang_code, text_dir=text_dir)

def write_footer(file_obj, fmt):
    core_write_footer(file_obj, fmt)

def save_pdf(path, content, large_print=False):
    return shared_save_pdf(
        path,
        content,
        large_print=large_print,
        fpdf_cls=FPDF,
        sanitize_latin1_fn=sanitize_latin1,
    )

def save_docx(path, content):
    return shared_save_docx(path, content, docx_module=docx)

def save_epub(path, title, content, lang_code="en", text_dir="ltr"):
    return shared_save_epub(
        path,
        title,
        content,
        lang_code=lang_code,
        text_dir=text_dir,
        epub_module=epub,
    )

def dispatch_save(cfg, path, memory, title):
    return shared_dispatch_save(
        cfg,
        path,
        memory,
        title,
        sanitize_model_output_fn=sanitize_model_output,
        apply_modern_punctuation_fn=apply_modern_punctuation,
        apply_modern_currency_fn=apply_modern_currency,
        apply_expanded_abbreviations_fn=apply_expanded_abbreviations,
        strip_synthetic_page_filename_headings_fn=strip_synthetic_page_filename_headings,
        get_output_lang_code_fn=get_output_lang_code,
        get_output_text_direction_fn=get_output_text_direction,
        save_docx_fn=save_docx,
        save_pdf_fn=save_pdf,
        save_epub_fn=save_epub,
    )

# --- AI API ENGINES ---
def handle_stream(response, output_path, fmt_type, file_obj, memory, log_cb, pause_cb=None):
    return shared_handle_stream(
        response,
        output_path,
        fmt_type,
        file_obj,
        memory,
        log_cb,
        pause_cb=pause_cb,
        heartbeat=heartbeat,
        sanitize_model_output_fn=sanitize_model_output,
        clean_text_fn=clean_text,
    )

def stream_with_cache(cache_key, request_fn, output_path, fmt, file_obj, memory, log_cb, pause_cb=None):
    return shared_stream_with_cache(
        cache_key,
        request_fn,
        output_path,
        fmt,
        file_obj,
        memory,
        log_cb,
        pause_cb=pause_cb,
        runtime=request_runtime,
        append_generated_text_fn=append_generated_text,
        handle_stream_fn=handle_stream,
    )

def generate_retry(client, model, contents, max_r=API_MAX_RETRIES, delay=API_BACKOFF_BASE_SEC, log_cb=print):
    return shared_generate_retry(
        client,
        model,
        contents,
        runtime=request_runtime,
        max_r=max_r,
        delay=delay,
        backoff_max_sec=API_BACKOFF_MAX_SEC,
        log_cb=log_cb,
    )

def build_payload(model, prompt, file_path=None, mime="image/png", file_bytes=None):
    return shared_build_payload(model, prompt, file_path=file_path, mime=mime, file_bytes=file_bytes)

# --- DOCUMENT PROCESSORS ---
def process_pdf(client, path, out, fmt, prompt, model, f_obj, mem, log_cb, confidence_cb=None, pause_cb=None, page_progress_cb=None, page_scope="", doc_profile="standard", auto_escalation_model=None):
    return shared_process_pdf(
        client,
        path,
        out,
        fmt,
        prompt,
        model,
        f_obj,
        mem,
        log_cb,
        confidence_cb=confidence_cb,
        pause_cb=pause_cb,
        page_progress_cb=page_progress_cb,
        page_scope=page_scope,
        doc_profile=doc_profile,
        auto_escalation_model=auto_escalation_model,
        script_dir=SCRIPT_DIR,
        pdf_reader_cls=PdfReader,
        pdf_writer_cls=PdfWriter,
        parse_pdf_page_scope_spec_fn=parse_pdf_page_scope_spec,
        normalize_pdf_page_scope_text_fn=normalize_pdf_page_scope_text,
        get_pdf_chunk_pages_fn=get_pdf_chunk_pages,
        sha256_file_fn=sha256_file,
        sha256_text_fn=sha256_text,
        build_payload_fn=build_payload,
        build_request_cache_key_fn=build_request_cache_key,
        stream_with_cache_fn=stream_with_cache,
        generate_retry_fn=generate_retry,
        split_text_for_fail_safe_fn=split_text_for_fail_safe,
        remove_fn=os.remove,
        exists_fn=os.path.exists,
    )

def process_text(client, path, out, ext, fmt, prompt, model, f_obj, mem, log_cb, pause_cb=None, page_progress_cb=None):
    return shared_process_text(
        client,
        path,
        out,
        ext,
        fmt,
        prompt,
        model,
        f_obj,
        mem,
        log_cb,
        pause_cb=pause_cb,
        text_chunk_chars=TEXT_CHUNK_CHARS,
        csv_to_accessible_text_fn=csv_to_accessible_text,
        clean_text_fn=clean_text,
        batch_text_chunks_fn=batch_text_chunks,
        build_request_cache_key_fn=build_request_cache_key,
        sha256_text_fn=sha256_text,
        stream_with_cache_fn=stream_with_cache,
        generate_retry_fn=generate_retry,
        docx_module=docx,
        openpyxl_module=openpyxl,
        subprocess_module=subprocess,
        page_progress_cb=page_progress_cb,
    )

def process_pptx(client, path, out, fmt, prompt, model, f_obj, mem, log_cb, pause_cb=None, page_progress_cb=None):
    return shared_process_pptx(
        client,
        path,
        out,
        fmt,
        prompt,
        model,
        f_obj,
        mem,
        log_cb,
        pause_cb=pause_cb,
        page_progress_cb=page_progress_cb,
        text_chunk_chars=TEXT_CHUNK_CHARS,
        clean_text_fn=clean_text,
        batch_text_chunks_fn=batch_text_chunks,
        build_request_cache_key_fn=build_request_cache_key,
        sha256_text_fn=sha256_text,
        stream_with_cache_fn=stream_with_cache,
        generate_retry_fn=generate_retry,
        subprocess_module=subprocess,
    )

def process_epub(client, path, out, fmt, prompt, model, f_obj, mem, log_cb, pause_cb=None, page_progress_cb=None):
    return shared_process_epub(
        client,
        path,
        out,
        fmt,
        prompt,
        model,
        f_obj,
        mem,
        log_cb,
        pause_cb=pause_cb,
        text_chunk_chars=TEXT_CHUNK_CHARS,
        clean_text_fn=clean_text,
        batch_text_chunks_fn=batch_text_chunks,
        build_request_cache_key_fn=build_request_cache_key,
        sha256_text_fn=sha256_text,
        stream_with_cache_fn=stream_with_cache,
        generate_retry_fn=generate_retry,
        epub_module=epub,
        page_progress_cb=page_progress_cb,
    )

def process_img(client, path, out, fmt, prompt, model, f_obj, mem, log_cb, pause_cb=None):
    return shared_process_img(
        client,
        path,
        out,
        fmt,
        prompt,
        model,
        f_obj,
        mem,
        log_cb,
        pause_cb=pause_cb,
        enhance_image_fn=enhance_image,
        build_payload_fn=build_payload,
        build_request_cache_key_fn=build_request_cache_key,
        sha256_file_fn=sha256_file,
        stream_with_cache_fn=stream_with_cache,
        generate_retry_fn=generate_retry,
        remove_fn=os.remove,
    )

class MainFrame(wx.Frame):
    def __init__(self):
        self.build_stamp = get_runtime_build_stamp()
        super().__init__(None, title=f"{APP_NAME} | Build {self.build_stamp}", size=(980, 760))
        self.autorun_spec_path = str(os.environ.get(AUTORUN_SPEC_ENV, "") or "").strip()
        self.autorun_active = bool(self.autorun_spec_path)
        self.autorun_close_on_finish = False
        self.cfg = load_json(CONFIG_FILE, {
            'format_type': 'html', 'model_name': 'gemini-2.5-flash', 'model_override': '', 'doc_profile': 'standard',
            'translate_mode': 'none', 'collision_mode': 'skip', 'merge_files': False,
            'translate_target': 'English',
            'image_descriptions': True, 'modernize_punctuation': False, 'unit_conversion': False,
            'abbrev_expansion': False, 'preserve_original_page_numbers': False, 'large_print': False, 'delete_source_on_success': False,
            'preserve_source_structure': True,
            'custom_prompt': '', 'custom_commands': '',
            'pdf_textlayer_audit': True,
            'page_confidence_scoring': False,
            'low_memory_mode': False,
            'memory_telemetry': False,
            'auto_save_processing_log': False,
            'academic_footnote_mode': 'endnotes',
            'academic_annotation_mode': 'inline',
            'output_dir': os.path.join(SCRIPT_DIR, "output_html"),
            'recursive_scan': False, 'dest_mode': 0, 'custom_dest': '',
            'mac_folder_access_confirmed': False,
            'scan_output_dir': os.path.join(SCRIPT_DIR, "Input_Scans"),
            'scan_dpi': 300,
            'scan_auto_start': False,
            'scan_extract_mode': 'manual',
            'scan_merge_before_queue': False,
            'scan_merge_extract_output': False,
            'scheduled_start_ts': None,
            'last_browse_dir': '',
            'risk_notice_acknowledged': False,
            'risk_notice_dont_show_again': False,
        })
        self.keys = load_json(KEYS_FILE, {})
        # One-time migration for legacy default config: TXT + output_txt => HTML default.
        if self.cfg.get('format_type') == 'txt' and str(self.cfg.get('output_dir', '')).endswith('output_txt'):
            self.cfg['format_type'] = 'html'
            self.cfg['output_dir'] = os.path.join(SCRIPT_DIR, "output_html")
            save_json(CONFIG_FILE, self.cfg)
        self.is_running = False
        self.is_paused = False
        self.stop_requested = False
        self.resume_incomplete_only = False
        self.current_processing_index = -1
        self.queue = []
        self.clients = {}
        self.license_status_text = "License status unavailable"
        self.license_validation = None
        self.license_public_key = None
        self.license_public_key_error = ""
        self.processing_log_lines = []
        self.total_pages_processed = 0
        self.current_file_page_total = 0
        self.current_file_page_done = 0
        self.current_file_unit_label = "items"
        self.current_file_ordinal = 0
        self.current_file_name = ""
        self.current_run_total_units = 0
        self.current_run_completed_units = 0
        self.current_run_started_ts = 0.0
        self._last_engine_activity_ts = 0.0
        self._last_progress_heartbeat_ts = 0.0
        self.pending_scan_merge_extract = False
        self.scheduled_start_ts = self._normalize_future_timestamp(self.cfg.get('scheduled_start_ts'))
        self.scheduled_start_triggered = False
        if self.scheduled_start_ts is None and self.cfg.get('scheduled_start_ts') is not None:
            self.cfg['scheduled_start_ts'] = None
            save_json(CONFIG_FILE, self.cfg)
        self._session_lock = threading.Lock()
        self.RefreshLicenseStatus()
        self.session_timer = wx.Timer(self)
        self.Bind(wx.EVT_TIMER, self.OnSessionCheckpoint, self.session_timer)
        self.session_timer.Start(2000)
        self.Bind(wx.EVT_CLOSE, self.OnMainClose)
        global ACTIVE_FRAME
        ACTIVE_FRAME = self
        self.InitUI()
        if self.autorun_active:
            wx.CallAfter(self.StartAutomationRunIfRequested)
        else:
            wx.CallAfter(self.EnsureRiskNoticeAcknowledged)
        if platform.system() == "Darwin" and not self.autorun_active:
            wx.CallAfter(self.EnsureMacFolderAccessOnLaunch)

    def ShowWelcome(self):
        dlg = wx.Dialog(self, title="Welcome to Chronicle", size=(620, 300))
        panel = wx.Panel(dlg)
        root = wx.BoxSizer(wx.VERTICAL)
        message = (
            "To use this app, you need API keys for the AI models of your choice:\n"
            "Google Gemini, Anthropic Claude, or OpenAI GPT.\n\n"
            "Select Open API Key Vault to enter keys now."
        )
        msg_ctrl = wx.StaticText(panel, label=message)
        msg_ctrl.Wrap(560)
        msg_ctrl.SetName("Welcome Message")
        root.Add(msg_ctrl, 1, wx.EXPAND | wx.ALL, 12)
        dlg.SetName(f"Welcome to Chronicle. {message}")

        buttons = wx.StdDialogButtonSizer()
        btn_ok = wx.Button(panel, wx.ID_OK, "Open API Key Vault")
        btn_cancel = wx.Button(panel, wx.ID_CANCEL, "Later")
        btn_ok.SetToolTip("Open the API key vault.")
        buttons.AddButton(btn_ok)
        buttons.AddButton(btn_cancel)
        buttons.Realize()
        root.Add(buttons, 0, wx.ALIGN_RIGHT | wx.ALL, 10)
        panel.SetSizer(root)

        wx.CallAfter(btn_ok.SetFocus)
        result = dlg.ShowModal()
        dlg.Destroy()
        if result == wx.ID_OK:
            self.OnApiKeys(None)

    def EnsureRiskNoticeAcknowledged(self):
        if FirstLaunchNoticeDialog is None:
            return True
        if self.cfg.get("risk_notice_acknowledged") and self.cfg.get("risk_notice_dont_show_again"):
            if not any(self.keys.values()):
                wx.CallAfter(self.ShowWelcome)
            return True
        dlg = FirstLaunchNoticeDialog(self)
        result = dlg.ShowModal()
        if result == wx.ID_OK:
            self.cfg["risk_notice_acknowledged"] = True
            self.cfg["risk_notice_dont_show_again"] = bool(dlg.ShouldSuppressFuturePrompts())
            save_json(CONFIG_FILE, self.cfg)
            dlg.Destroy()
            if not any(self.keys.values()):
                wx.CallAfter(self.ShowWelcome)
            return True
        dlg.Destroy()
        wx.CallAfter(self.Close)
        return False

    def _session_payload(self):
        return shared_build_session_payload(
            SESSION_SCHEMA_VERSION,
            self.cfg,
            self.queue,
            self.is_running,
            self.is_paused,
        )

    def SaveActiveSession(self):
        payload = self._session_payload()
        try:
            shared_save_active_session_file(SESSION_FILE, payload, session_lock=self._session_lock)
        except Exception as ex:
            self.Log(f"Warning: could not save active session ({ex})", engine_event=True)

    def DeleteActiveSession(self):
        try:
            shared_delete_active_session_file(SESSION_FILE)
        except Exception as ex:
            self.Log(f"Warning: could not remove active session file ({ex})", engine_event=True)

    def OnSessionCheckpoint(self, event):
        self.CheckScheduledStart()
        if self.is_running or self._has_incomplete_items(self.queue):
            self.SaveActiveSession()

    def _normalize_future_timestamp(self, value):
        return shared_normalize_future_timestamp(value)

    def _format_timestamp_local(self, ts):
        return shared_format_timestamp_local(ts)

    def _set_scheduled_start(self, ts):
        normalized = self._normalize_future_timestamp(ts)
        self.scheduled_start_ts = normalized
        self.cfg['scheduled_start_ts'] = normalized
        save_json(CONFIG_FILE, self.cfg)
        self.SaveActiveSession()
        if hasattr(self, "schedule_summary"):
            self.UpdateScheduleSummary()

    def UpdateScheduleSummary(self):
        label = shared_build_schedule_summary_label(self.scheduled_start_ts)
        self.schedule_summary.SetLabel(label)
        self.schedule_summary.SetName(f"Scheduled Extraction Status ({label})")

    def CheckScheduledStart(self):
        if not shared_should_trigger_scheduled_start(
            self.scheduled_start_ts,
            is_running=self.is_running,
            scheduled_start_triggered=self.scheduled_start_triggered,
        ):
            return
        self.scheduled_start_triggered = True
        scheduled_at = self.scheduled_start_ts
        self._set_scheduled_start(None)
        if not self.queue:
            self.Log(
                f"Scheduled start reached at {self._format_timestamp_local(scheduled_at)}, but the queue is empty. Schedule cleared.",
                engine_event=True,
            )
            self.scheduled_start_triggered = False
            return
        self.Log(
            f"Scheduled start reached at {self._format_timestamp_local(scheduled_at)}. Starting extraction.",
            engine_event=True,
        )
        self.OnStart(None)
        self.scheduled_start_triggered = False

    def _has_incomplete_items(self, queue_rows):
        return shared_has_incomplete_items(queue_rows, SESSION_TERMINAL_STATUSES)

    def _sync_controls_from_cfg(self):
        fmt_selection = {'html': 0, 'txt': 1, 'docx': 2, 'md': 3, 'pdf': 4, 'json': 5, 'csv': 6, 'epub': 7}
        self.format_choice.SetSelection(fmt_selection.get(self.cfg.get('format_type', 'html'), 0))
        tmap = {'none': 0, 'both': 1, 'english_only': 2}
        self.run_translate_choice.SetSelection(tmap.get(self.cfg.get('translate_mode', 'none'), 0))
        self.run_punct_choice.SetSelection(1 if self.cfg.get('modernize_punctuation', False) else 0)
        self.run_units_choice.SetSelection(1 if self.cfg.get('unit_conversion', False) else 0)
        self.run_abbrev_choice.SetSelection(1 if self.cfg.get('abbrev_expansion', False) else 0)
        self.run_images_choice.SetSelection(0 if self.cfg.get('image_descriptions', True) else 1)
        self.run_page_numbers_choice.SetSelection(1 if self.cfg.get('preserve_original_page_numbers', False) else 0)
        self.run_merge_choice.SetSelection(1 if self.cfg.get('merge_files', False) else 0)
        self.run_largeprint_choice.SetSelection(1 if self.cfg.get('large_print', False) else 0)
        self.SetPdfPageScopeSelection(self.cfg.get('pdf_page_scope', ''))
        self.choice_recursive.SetSelection(1 if bool(self.cfg.get('recursive_scan', False)) else 0)
        self.dest_choice.SetSelection(1 if int(self.cfg.get('dest_mode', 0)) == 1 else 0)
        self.txt_dest.SetValue(self.cfg.get('custom_dest', ''))
        self.chk_preserve_structure.SetValue(bool(self.cfg.get('preserve_source_structure', True)))
        self.chk_delete_originals.SetValue(bool(self.cfg.get('delete_source_on_success', False)))
        selected_profile = self.cfg.get('doc_profile', 'standard')
        profile_index = next((i for i, (k, _) in enumerate(PROFILE_CHOICES) if k == selected_profile), 0)
        self.choice_profile.SetSelection(profile_index)
        self.scheduled_start_ts = self._normalize_future_timestamp(self.cfg.get('scheduled_start_ts'))
        if self.scheduled_start_ts is None and self.cfg.get('scheduled_start_ts') is not None:
            self.cfg['scheduled_start_ts'] = None
            save_json(CONFIG_FILE, self.cfg)
        self.UpdateModelGuide()
        self.UpdateProfileGuide()
        self.UpdatePreflightSummary()
        self.UpdateLargePrintVisibility()
        self.UpdateScheduleSummary()
        self.OnDestModeChanged()

    def CheckForRecoverableSession(self):
        if not os.path.exists(SESSION_FILE):
            return
        session = load_json(SESSION_FILE, {})
        if not isinstance(session, dict):
            return
        queue_rows = session.get("queue", [])
        cfg = session.get("cfg", {})
        if not isinstance(queue_rows, list) or not isinstance(cfg, dict):
            return
        if not self._has_incomplete_items(queue_rows):
            return

        dlg = SessionRecoveryDialog(self)
        choice = dlg.ShowModal()
        dlg.Destroy()

        if choice == SessionRecoveryDialog.START_FRESH:
            self.DeleteActiveSession()
            self.queue = []
            self.RefreshQueue()
            self.SafeSetStatusText("Started with a fresh session.")
            return

        self.cfg.update(cfg)
        self.queue = shared_restore_session_queue(
            queue_rows,
            self.cfg,
            row_setting_keys=ROW_SETTING_KEYS,
            label_from_model_fn=self.LabelFromModel,
        )

        self._sync_controls_from_cfg()
        self.RefreshQueue()
        self.SaveActiveSession()

        if choice == SessionRecoveryDialog.RESUME:
            self.resume_incomplete_only = True
            wx.CallAfter(self.OnStart, None)
        else:
            self.SafeSetStatusText("Session loaded. Review queue and press Start Extraction when ready.")

    def WaitWhilePaused(self):
        return shared_wait_while_paused(
            is_running=self.is_running,
            is_paused=self.is_paused,
            stop_requested=self.stop_requested,
            heartbeat_ping=heartbeat.ping,
            sleep_fn=time.sleep,
            stop_requested_error_cls=StopRequestedError,
        )

    def OnAppQuitRequested(self, _event):
        self.Close()

    def OnMainClose(self, event):
        has_incomplete = self._has_incomplete_items(self.queue)
        if self.is_running or has_incomplete:
            active_count = (
                shared_count_active_queue_items(
                    self.queue,
                    terminal_statuses=SESSION_TERMINAL_STATUSES,
                )
                if self.is_running
                else shared_count_saved_queue_items(
                    self.queue,
                    terminal_statuses=SESSION_TERMINAL_STATUSES,
                )
            )
            dlg = CloseRunningDialog(self, active_count, is_running=self.is_running)
            choice = dlg.ShowModal()
            dlg.Destroy()
            close_plan = shared_prepare_running_close(
                choice,
                keep_open_value=CloseRunningDialog.KEEP_OPEN,
                save_exit_value=CloseRunningDialog.SAVE_EXIT,
                discard_exit_value=CloseRunningDialog.DISCARD_EXIT,
                is_running=self.is_running,
            )
            if close_plan["pause_run"]:
                self.is_paused = True
            if close_plan["should_veto"]:
                if close_plan["pause_run"] and 0 <= self.current_processing_index < len(self.queue):
                    self.SetQueueStatus(self.current_processing_index, "Paused")
                if close_plan["pause_run"]:
                    self.SaveActiveSession()
                event.Veto()
                self.file_list.SetFocus()
                return
            if close_plan["pause_current_row_directly"]:
                shared_pause_current_processing_row(self.queue, self.current_processing_index)
            if close_plan["discard_session"]:
                self.stop_requested = True
                for row in self.queue:
                    if str(row.get("status", "Queued")) not in SESSION_TERMINAL_STATUSES:
                        row["status"] = "Stopped"
                self.DeleteActiveSession()
            elif close_plan["save_session"]:
                self.SaveActiveSession()

        self.session_timer.Stop()
        if self._has_incomplete_items(self.queue):
            self.SaveActiveSession()
        else:
            self.DeleteActiveSession()
        event.Skip()

    def EnsureMacFolderAccessOnLaunch(self):
        if self.cfg.get('mac_folder_access_confirmed', False):
            return
        self.PromptMacFolderAccess(force=False)

    def PromptMacFolderAccess(self, force=False):
        if platform.system() != "Darwin":
            return
        if self.cfg.get('mac_folder_access_confirmed', False) and not force:
            return

        msg = (
            "Chronicle may need macOS permission to access Documents and subfolders.\n\n"
            "Click Yes to grant access by selecting your Documents folder now."
        )
        res = wx.MessageBox(msg, "Grant Mac Folder Access", wx.YES_NO | wx.ICON_QUESTION)
        if res != wx.YES:
            return

        docs_dir = os.path.expanduser("~/Documents")
        dlg = wx.DirDialog(
            self,
            "Select Documents folder to grant Chronicle access",
            defaultPath=docs_dir,
            style=wx.DD_DEFAULT_STYLE | wx.DD_DIR_MUST_EXIST,
        )
        selected = None
        if dlg.ShowModal() == wx.ID_OK:
            selected = dlg.GetPath()
        dlg.Destroy()
        if not selected:
            return

        try:
            os.listdir(selected)
            self.cfg['mac_folder_access_confirmed'] = True
            save_json(CONFIG_FILE, self.cfg)
            self.Log(f"macOS folder access confirmed: {selected}")
        except Exception as ex:
            wx.MessageBox(
                "Access was not granted.\n\n"
                "Please allow Chronicle in System Settings > Privacy & Security > Files and Folders, "
                "or run File > Check Mac Folder Access again.\n\n"
                f"Details: {ex}",
                "Mac Permission Needed",
                wx.OK | wx.ICON_WARNING,
            )

    def InitUI(self):
        def set_button_icon(btn, art_id):
            try:
                bmp = wx.ArtProvider.GetBitmap(art_id, wx.ART_BUTTON, (16, 16))
                if bmp and bmp.IsOk():
                    btn.SetBitmap(bmp)
            except Exception:
                pass

        ChronicleMenuBar(self).install()

        panel = wx.Panel(self)
        root = wx.BoxSizer(wx.VERTICAL)

        self.queue_panel = QueuePanel(panel, self, empty_placeholder=QUEUE_EMPTY_PLACEHOLDER)
        self.file_list = self.queue_panel.file_list
        root.Add(self.queue_panel, 1, wx.EXPAND)

        queue_button_bindings = {
            'add_files': (wx.EVT_BUTTON, self.OnAddFiles),
            'add_folder': (wx.EVT_BUTTON, self.OnAddFolder),
            'remove_selected': (wx.EVT_BUTTON, self.OnRemoveSelected),
            'task_actions': (wx.EVT_BUTTON, self.OnTaskActionsButton),
            'select_all': (wx.EVT_BUTTON, self.OnSelectAllQueue),
            'deselect_all': (wx.EVT_BUTTON, self.OnDeselectAllQueue),
            'clear_list': (wx.EVT_BUTTON, self.OnClearQueue),
        }
        queue_section = build_queue_action_section(
            panel,
            bind=lambda widget, name: bind_named(widget, name, queue_button_bindings),
            set_button_icon=set_button_icon,
        )
        self.btn_add = queue_section['btn_add']
        self.btn_add_folder = queue_section['btn_add_folder']
        self.btn_remove = queue_section['btn_remove']
        self.btn_task_actions = queue_section['btn_task_actions']
        self.btn_select_all = queue_section['btn_select_all']
        self.btn_deselect_all = queue_section['btn_deselect_all']
        self.btn_clear = queue_section['btn_clear']
        root.Add(queue_section['queue_actions_label'], 0, wx.LEFT | wx.RIGHT, 10)
        root.Add(queue_section['queue_actions'], 0, wx.EXPAND)
        root.Add(queue_section['queue_tools'], 0, wx.EXPAND)

        settings_bindings = {
            'format_choice': (wx.EVT_CHOICE, self.OnFormatChoiceChanged),
            'apply_settings': (wx.EVT_BUTTON, self.OnApplySettingsToQueue),
            'run_preflight': (wx.EVT_BUTTON, self.OnRunPreflight),
            'pdf_pages': (wx.EVT_CHOICE, self.OnPdfPageScopeChanged),
            'run_option': (wx.EVT_CHOICE, self.OnRunOptionChoiceChanged),
            'dest_mode': (wx.EVT_CHOICE, self.OnDestModeChanged),
            'choose_dest': (wx.EVT_BUTTON, self.OnChooseDest),
            'profile_choice': (wx.EVT_CHOICE, self.OnProfileChoiceChanged),
            'safety_option': (wx.EVT_CHECKBOX, self.OnSafetyOptionsChanged),
            'start': (wx.EVT_BUTTON, self.OnStart),
            'schedule': (wx.EVT_BUTTON, self.OnScheduleExtraction),
        }
        settings_section = build_settings_section(
            panel,
            cfg=self.cfg,
            profile_choices=PROFILE_CHOICES,
            get_pdf_page_items_fn=self.GetPdfPageScopeChoiceItems,
            bind=lambda widget, name: bind_named(widget, name, settings_bindings),
        )
        self.format_choice = settings_section['format_choice']
        self.btn_apply_settings = settings_section['btn_apply_settings']
        self.btn_run_preflight = settings_section['btn_run_preflight']
        self.preflight_summary = settings_section['preflight_summary']
        self.run_translate_choice = settings_section['run_translate_choice']
        self.run_punct_choice = settings_section['run_punct_choice']
        self.run_units_choice = settings_section['run_units_choice']
        self.run_abbrev_choice = settings_section['run_abbrev_choice']
        self.run_images_choice = settings_section['run_images_choice']
        self.run_page_numbers_choice = settings_section['run_page_numbers_choice']
        self.run_merge_choice = settings_section['run_merge_choice']
        self.lbl_run_largeprint = settings_section['lbl_run_largeprint']
        self.run_largeprint_choice = settings_section['run_largeprint_choice']
        self.lbl_pdf_pages = settings_section['lbl_pdf_pages']
        self.choice_pdf_pages = settings_section['choice_pdf_pages']
        self.choice_recursive = settings_section['choice_recursive']
        self.dest_choice = settings_section['dest_choice']
        self.txt_dest = settings_section['txt_dest']
        self.btn_dest = settings_section['btn_dest']
        self.choice_profile = settings_section['choice_profile']
        self.chk_preserve_structure = settings_section['chk_preserve_structure']
        self.chk_delete_originals = settings_section['chk_delete_originals']
        self.btn_start = settings_section['btn_start']
        self.btn_schedule = settings_section['btn_schedule']
        set_button_icon(self.btn_start, wx.ART_GO_FORWARD)
        set_button_icon(self.btn_schedule, wx.ART_TIP)

        root.Add(settings_section['run_settings_label'], 0, wx.LEFT | wx.RIGHT | wx.TOP, 10)
        root.Add(settings_section['top'], 0, wx.EXPAND | wx.LEFT | wx.RIGHT | wx.BOTTOM, 10)

        self.UpdateModelGuide()

        root.Add(settings_section['runtime_opts'], 0, wx.EXPAND | wx.LEFT | wx.RIGHT | wx.BOTTOM, 10)
        self.SetPdfPageScopeSelection(self.cfg.get('pdf_page_scope', ''))
        root.Add(settings_section['opts'], 0, wx.EXPAND | wx.LEFT | wx.RIGHT | wx.BOTTOM, 10)
        root.Add(settings_section['safety_opts'], 0, wx.EXPAND | wx.LEFT | wx.RIGHT | wx.BOTTOM, 10)
        root.Add(settings_section['apply_settings_row'], 0, wx.LEFT | wx.RIGHT | wx.BOTTOM, 10)
        root.Add(settings_section['preflight_row'], 0, wx.EXPAND | wx.LEFT | wx.RIGHT | wx.BOTTOM, 10)
        root.Add(settings_section['run_actions'], 0, wx.EXPAND | wx.LEFT | wx.RIGHT | wx.TOP, 10)

        self.schedule_summary = wx.StaticText(panel, label="No extraction scheduled.")
        self.schedule_summary.SetName("Scheduled Extraction Status")
        root.Add(self.schedule_summary, 0, wx.LEFT | wx.RIGHT | wx.BOTTOM, 10)
        self.UpdateScheduleSummary()

        progress_box, self.progress_gauge, self.progress_summary = build_progress_section(panel)
        root.Add(progress_box, 0, wx.EXPAND | wx.LEFT | wx.RIGHT, 10)

        log_label, self.lt, log_actions, self.btn_save_log = build_log_section(
            panel,
            save_log_handler=self.OnSaveLog,
        )
        root.Add(log_label, 0, wx.LEFT | wx.RIGHT | wx.TOP, 10)
        root.Add(self.lt, 1, wx.EXPAND | wx.ALL, 10)
        root.Add(log_actions, 0, wx.LEFT | wx.RIGHT | wx.BOTTOM, 10)
        panel.SetSizer(root)
        self.CreateStatusBar(1)
        self.SafeSetStatusText(f"Ready. Build {self.build_stamp}")
        self.UpdateProfileGuide()
        self.UpdateLargePrintVisibility()
        self.OnDestModeChanged()
        self.UpdateQueueButtons()
        self.UpdateProgressIndicators(force_status=True)
        self.file_list.SetFocus()
        if not self.autorun_active:
            wx.CallAfter(self.CheckForRecoverableSession)

    def StartAutomationRunIfRequested(self):
        if not self.autorun_active:
            return
        try:
            with open(self.autorun_spec_path, "r", encoding="utf-8") as fh:
                spec = json.load(fh)
            if not isinstance(spec, dict):
                raise ValueError("Autorun spec must be a JSON object.")

            raw_paths = spec.get("paths", [])
            if isinstance(raw_paths, str):
                raw_paths = [raw_paths]
            paths = [os.path.abspath(str(path)) for path in raw_paths if str(path or "").strip()]
            if not paths:
                raise ValueError("Autorun spec must include at least one path.")

            for path in paths:
                if not os.path.exists(path):
                    raise FileNotFoundError(path)

            output_dir = str(spec.get("output_dir") or spec.get("custom_dest") or "").strip()
            profile_key = str(spec.get("doc_profile") or "standard").strip() or "standard"
            format_type = str(spec.get("format_type") or "html").strip() or "html"
            page_scope = normalize_pdf_page_scope_text(
                spec.get("pdf_page_scope") or spec.get("page_scope") or ""
            )

            self.autorun_close_on_finish = bool(spec.get("close_on_finish", True))
            self.cfg["doc_profile"] = profile_key
            self.cfg["format_type"] = format_type
            self.cfg["pdf_page_scope"] = page_scope
            self.cfg["merge_files"] = bool(spec.get("merge_files", False))
            self.cfg["preserve_source_structure"] = bool(spec.get("preserve_source_structure", True))
            self.cfg["delete_source_on_success"] = bool(spec.get("delete_source_on_success", False))
            self.cfg["auto_save_processing_log"] = bool(spec.get("auto_save_processing_log", True))
            self.cfg["modernize_punctuation"] = bool(spec.get("modernize_punctuation", self.cfg.get("modernize_punctuation", False)))
            self.cfg["unit_conversion"] = bool(spec.get("unit_conversion", self.cfg.get("unit_conversion", False)))
            self.cfg["abbrev_expansion"] = bool(spec.get("abbrev_expansion", self.cfg.get("abbrev_expansion", False)))
            self.cfg["image_descriptions"] = bool(spec.get("image_descriptions", self.cfg.get("image_descriptions", True)))
            self.cfg["mac_folder_access_confirmed"] = True
            if output_dir:
                self.cfg["dest_mode"] = 1
                self.cfg["custom_dest"] = output_dir
                os.makedirs(output_dir, exist_ok=True)
            else:
                self.cfg["dest_mode"] = 0
                self.cfg["custom_dest"] = ""

            requested_model = str(spec.get("model_name") or "").strip()
            if requested_model:
                self.cfg["model_name"] = self.ResolveModelNameForAvailableKeys(requested_model)
            else:
                self.cfg["model_name"] = self.ResolveModelNameForAvailableKeys(
                    self.GetPreferredModelName(profile_key)
                )

            self._sync_controls_from_cfg()
            self.queue = []
            settings = self.BuildSettingsFromControls()
            shared_add_path_entries(
                self.queue,
                paths,
                settings=settings,
                engine_label=self.LabelFromModel(settings.get("model_name", "gemini-2.5-flash")),
                row_setting_keys=ROW_SETTING_KEYS,
                source_root=os.path.dirname(paths[0]) if len(paths) == 1 else None,
            )
            self.RefreshQueue()
            self.SaveActiveSession()
            self.Log(
                f"[Autorun] Loaded {len(paths)} file(s) from {self.autorun_spec_path} via packaged app.",
                engine_event=True,
            )
            self.OnStart(None)
        except Exception as ex:
            self.Log(f"[Autorun] Failed to start automation run: {ex}", engine_event=True)
            if self.autorun_close_on_finish:
                wx.CallAfter(self.Close)

    def ModelFromLabel(self, label):
        return {
            'Gemini 2.5 Flash': 'gemini-2.5-flash',
            'Gemini 2.5 Pro': 'gemini-2.5-pro',
            'Claude Sonnet 4': DEFAULT_CLAUDE_MODEL,
            'GPT-4o': 'gpt-4o'
        }.get(label, 'gemini-2.5-flash')

    def LabelFromModel(self, model_name):
        return {
            'gemini-2.5-flash': 'Gemini 2.5 Flash',
            'gemini-2.5-pro': 'Gemini 2.5 Pro',
            DEFAULT_CLAUDE_MODEL: 'Claude Sonnet 4',
            'gpt-4o': 'GPT-4o',
        }.get(model_name, 'Gemini 2.5 Flash')

    def GetModelGuideText(self, label):
        guides = {
            'Automatic': "Chronicle chooses the engine from the preset and a quick document preflight, with bounded PDF escalation when needed.",
            'Gemini 2.5 Flash': "Best for speed and lower cost on large batches. Good default for most standard extraction jobs.",
            'Gemini 2.5 Pro': "Best for difficult scans, nuanced documents, and maximum extraction fidelity. Slower and higher cost than Flash.",
            'Claude Sonnet 4': "Best for strong structured writing quality and complex instruction-following. Good for polished narrative outputs.",
            'GPT-4o': "Best for broad multimodal reasoning and mixed document/image tasks. Strong general-purpose fallback across formats.",
        }
        return guides.get(label, guides['Gemini 2.5 Flash'])

    def SafeSetStatusText(self, text):
        status_bar = self.GetStatusBar()
        if status_bar is not None:
            self.SetStatusText(text)

    def HasProviderKey(self, vendor):
        vendor_key = "claude" if vendor == "anthropic" else str(vendor or "")
        return bool(self.keys.get(vendor_key))

    def ResolveModelNameForAvailableKeys(self, model_name):
        return shared_resolve_model_for_available_keys(
            model_name,
            has_vendor_key_fn=self.HasProviderKey,
        )

    def GetPreferredModelName(self, profile_key):
        return get_preferred_profile_model(profile_key, self.cfg)

    def UpdateModelGuide(self):
        if hasattr(self, 'choice_profile'):
            profile_key = PROFILE_CHOICES[self.choice_profile.GetSelection()][0]
        else:
            profile_key = self.cfg.get('doc_profile', 'standard')
        preferred_model_name = self.GetPreferredModelName(profile_key)
        model_name = self.ResolveModelNameForAvailableKeys(preferred_model_name)
        label = self.LabelFromModel(model_name)
        manual_override = str(self.cfg.get("model_override", "") or "").strip()
        model_help = self.GetModelGuideText('Automatic' if not manual_override else label)
        speed_warning = get_processing_speed_warning(profile_key, model_name)
        status_text = f"{('Automatic' if not manual_override else label)}: {model_help}"
        if manual_override:
            status_text = f"{status_text} Preference override is set to {self.LabelFromModel(preferred_model_name)}."
        else:
            status_text = f"{status_text} Baseline recommendation is {label} for this preset."
        if model_name != preferred_model_name and any(self.keys.values()):
            preferred_label = self.LabelFromModel(preferred_model_name)
            status_text = f"{status_text} Using {label} because {preferred_label} needs an API key that is not configured."
        if speed_warning:
            status_text = f"{status_text} {speed_warning}"
        status_bar = self.GetStatusBar()
        if status_bar is not None:
            self.SafeSetStatusText(status_text[:180])

    def GetPreflightTargetIndex(self):
        selected = self.GetSelectedIndices()
        if selected:
            return selected[0]
        if self.queue:
            return 0
        return -1

    def BuildPreflightReport(self, path, settings):
        ext = os.path.splitext(path)[1].lower()
        preferred_model_name = str(settings.get("model_name") or self.GetPreferredModelName(settings.get("doc_profile", "standard")))
        routing = select_execution_model_for_job(
            path,
            ext,
            settings,
            preferred_model_name,
            pdf_reader_factory=PdfReader,
            normalize_pdf_page_scope_text_fn=normalize_pdf_page_scope_text,
            parse_pdf_page_scope_spec_fn=parse_pdf_page_scope_spec,
        )
        selected_model = self.ResolveModelNameForAvailableKeys(str(routing.get("model_name") or preferred_model_name))
        profile_label = PROFILE_KEY_TO_LABEL.get(settings.get("doc_profile", "standard"), "Standard")
        unit_estimate = shared_estimate_current_file_total_units(
            ext,
            path,
            settings,
            pdf_reader_factory=PdfReader,
            normalize_pdf_page_scope_text_fn=normalize_pdf_page_scope_text,
            parse_pdf_page_scope_spec_fn=parse_pdf_page_scope_spec,
            pptx_slide_count_fn=lambda pptx_path: len(__import__('pptx').Presentation(pptx_path).slides),
            estimate_text_work_units_fn=lambda source_path, source_ext, job_cfg: shared_estimate_text_work_units(
                source_path,
                source_ext,
                text_chunk_chars=TEXT_CHUNK_CHARS,
                csv_to_accessible_text_fn=csv_to_accessible_text,
                clean_text_fn=clean_text,
                batch_text_chunks_fn=batch_text_chunks,
                docx_module=docx,
                openpyxl_module=openpyxl,
            ),
        )
        file_name = os.path.basename(path)
        mode_label = "manual override" if str(settings.get("model_override", "") or "").strip() else "automatic preflight"
        strategy = "Deep scan path" if ext in {".pdf", ".jpg", ".jpeg", ".png", ".bmp", ".tiff", ".tif", ".webp"} else "Direct source-text path"
        reason = str(routing.get("routing_reason", "") or "No preflight reason available.")
        if ext in [".docx", ".txt", ".md", ".rtf", ".csv", ".js", ".xlsx", ".xls", ".epub", ".pptx", ".ppt"]:
            reason = (
                "Strong source text is available directly from the file, so Chronicle will replicate that text first "
                "and only apply conservative accessibility repair."
            )
        elif "text layer" in reason.lower() or "born-digital" in reason.lower():
            strategy = "Direct text-first PDF path"
        scope_detail = ""
        if unit_estimate.get("selected_scope"):
            scope_detail = (
                f" Scope: pages {unit_estimate['selected_scope']} "
                f"({unit_estimate['selected_count']} selected of {unit_estimate['source_total']})."
            )
        elif unit_estimate.get("total_units", 0):
            scope_detail = f" Estimated work: {unit_estimate['total_units']} {unit_estimate.get('unit_label', 'unit(s)')}."
        summary = (
            f"Preflight for {file_name}: {profile_label}, {mode_label}. "
            f"Engine: {self.LabelFromModel(selected_model)}. Strategy: {strategy}. {reason}{scope_detail}"
        )
        details = "\n".join(
            [
                f"File: {file_name}",
                f"Preset: {profile_label}",
                f"Mode: {mode_label}",
                f"Chosen engine: {self.LabelFromModel(selected_model)}",
                f"Strategy: {strategy}",
                f"Reason: {reason}",
                (
                    f"Selected PDF scope: {unit_estimate['selected_scope']} "
                    f"({unit_estimate['selected_count']} of {unit_estimate['source_total']} pages)"
                    if unit_estimate.get("selected_scope")
                    else f"Estimated work units: {unit_estimate.get('total_units', 0)} {unit_estimate.get('unit_label', 'unit(s)')}"
                ),
            ]
        )
        return {
            "summary": summary,
            "details": details,
            "routing": routing,
            "selected_model": selected_model,
        }

    def UpdatePreflightSummary(self):
        if not hasattr(self, "preflight_summary"):
            return
        idx = self.GetPreflightTargetIndex()
        if idx < 0 or idx >= len(self.queue):
            self.preflight_summary.SetLabel(
                "Preflight ready. Select a queued file and run preflight to see what Chronicle detects before extraction starts."
            )
            return
        try:
            settings = self.BuildSettingsFromControls()
            report = self.BuildPreflightReport(self.queue[idx]["path"], settings)
            self.preflight_summary.SetLabel(report["summary"])
        except Exception as ex:
            self.preflight_summary.SetLabel(f"Preflight ready for {os.path.basename(self.queue[idx]['path'])}. Chronicle could not inspect it yet ({ex}).")

    def OnRunPreflight(self, event):
        idx = self.GetPreflightTargetIndex()
        if idx < 0 or idx >= len(self.queue):
            wx.MessageBox(
                "Add a file to the queue first, then select it and run preflight.",
                "No File To Preflight",
                wx.OK | wx.ICON_INFORMATION,
            )
            if event:
                event.Skip()
            return
        try:
            settings = self.BuildSettingsFromControls()
            report = self.BuildPreflightReport(self.queue[idx]["path"], settings)
            self.preflight_summary.SetLabel(report["summary"])
            self.Log(f"[Preflight] {report['summary']}", engine_event=True)
            wx.MessageBox(report["details"], "Document Preflight", wx.OK | wx.ICON_INFORMATION)
        except Exception as ex:
            wx.MessageBox(
                f"Chronicle could not complete document preflight.\n\nDetails: {ex}",
                "Preflight Error",
                wx.OK | wx.ICON_WARNING,
            )
        if event:
            event.Skip()

    def GetPdfPageScopeChoiceItems(self, selected_scope=""):
        selected_label = pdf_page_scope_choice_label_from_value(selected_scope)
        items = ["All Pages", "Page 1", "Pages 1-5"]
        if selected_label not in items and selected_label:
            items.append(selected_label)
        items.append("Custom...")
        return items

    def GetSelectedPdfPageScope(self):
        if not hasattr(self, "choice_pdf_pages"):
            return ""
        selection = self.choice_pdf_pages.GetStringSelection().strip()
        return normalize_pdf_page_scope_text(pdf_page_scope_value_from_choice_label(selection))

    def SetPdfPageScopeSelection(self, scope_text):
        normalized = normalize_pdf_page_scope_text(scope_text)
        if not hasattr(self, "choice_pdf_pages"):
            return
        items = self.GetPdfPageScopeChoiceItems(normalized)
        self.choice_pdf_pages.SetItems(items)
        target_label = pdf_page_scope_choice_label_from_value(normalized)
        if target_label not in items:
            target_label = "All Pages"
        self.choice_pdf_pages.SetSelection(items.index(target_label))

    def PromptForCustomPdfPageScope(self):
        current_scope = self.GetSelectedPdfPageScope()
        if current_scope in {"", "1", "1-5"}:
            current_scope = ""
        dialog = wx.TextEntryDialog(
            self,
            "Enter single pages or ranges.\n"
            "Use spaces between groups and hyphens for ranges.\n"
            "Examples: 7 or 1-30 or 1-30 185-220 530-574",
            "Custom PDF Pages",
            value=current_scope,
        )
        try:
            while True:
                if dialog.ShowModal() != wx.ID_OK:
                    return None
                custom_scope = normalize_pdf_page_scope_text(dialog.GetValue())
                if is_valid_pdf_page_scope_text(custom_scope):
                    return custom_scope
                wx.MessageBox(
                    "Use spaces between groups and hyphens for ranges.\n\n"
                    "Valid examples:\n"
                    "7\n"
                    "1-30\n"
                    "1-30 185-220 530-574\n\n"
                    "Commas, semicolons, and new lines also work if you prefer them.",
                    "Invalid PDF Page Range",
                    wx.OK | wx.ICON_WARNING,
                )
                dialog.SetValue(custom_scope)
        finally:
            dialog.Destroy()

    def ApplyProfilePresetToControls(self, profile_key):
        control_settings = build_control_settings_for_profile(
            self.cfg,
            profile_key,
            selected_model_name=self.cfg.get("model_name", "gemini-2.5-flash"),
        )
        tmap = {'none': 0, 'both': 1, 'english_only': 2}
        self.run_translate_choice.SetSelection(tmap.get(control_settings.get("translate_mode", "none"), 0))
        self.run_punct_choice.SetSelection(1 if control_settings.get("modernize_punctuation", False) else 0)
        self.run_units_choice.SetSelection(1 if control_settings.get("unit_conversion", False) else 0)
        self.run_abbrev_choice.SetSelection(1 if control_settings.get("abbrev_expansion", False) else 0)
        self.run_images_choice.SetSelection(0 if control_settings.get("image_descriptions", True) else 1)
        self.run_page_numbers_choice.SetSelection(1 if control_settings.get("preserve_original_page_numbers", False) else 0)
        self.run_merge_choice.SetSelection(1 if control_settings.get("merge_files", False) else 0)
        self.run_largeprint_choice.SetSelection(1 if control_settings.get("large_print", False) else 0)
        self.UpdateLargePrintVisibility()

    def BuildSettingsFromControls(self):
        fmt_map = {0: 'html', 1: 'txt', 2: 'docx', 3: 'md', 4: 'pdf', 5: 'json', 6: 'csv', 7: 'epub'}
        t_rev = {0: 'none', 1: 'both', 2: 'english_only'}
        selected_profile = PROFILE_CHOICES[self.choice_profile.GetSelection()][0]
        preferred_model_name = self.GetPreferredModelName(selected_profile)
        selected_model_name = self.ResolveModelNameForAvailableKeys(preferred_model_name)
        cfg = dict(self.cfg)
        cfg['format_type'] = fmt_map.get(self.format_choice.GetSelection(), 'html')
        apply_profile_preset(
            cfg,
            selected_profile,
            selected_model_name=selected_model_name,
            keep_selected_model=True,
        )
        cfg['model_name'] = selected_model_name
        cfg['translate_mode'] = t_rev.get(self.run_translate_choice.GetSelection(), 'none')
        cfg['modernize_punctuation'] = self.run_punct_choice.GetSelection() == 1
        cfg['unit_conversion'] = self.run_units_choice.GetSelection() == 1
        cfg['abbrev_expansion'] = self.run_abbrev_choice.GetSelection() == 1
        cfg['image_descriptions'] = self.run_images_choice.GetSelection() == 0
        cfg['preserve_original_page_numbers'] = self.run_page_numbers_choice.GetSelection() == 1
        cfg['merge_files'] = self.run_merge_choice.GetSelection() == 1
        cfg['large_print'] = self.run_largeprint_choice.GetSelection() == 1
        cfg['pdf_page_scope'] = self.GetSelectedPdfPageScope()
        return {key: cfg.get(key) for key in ROW_SETTING_KEYS}

    def NormalizeRowSettings(self, row):
        current = row.get("settings") if isinstance(row.get("settings"), dict) else {}
        normalized = dict(current)
        defaults = self.BuildSettingsFromControls()
        for key, val in defaults.items():
            if key not in normalized:
                normalized[key] = val
        normalized["model_name"] = self.ResolveModelNameForAvailableKeys(
            normalized.get("model_name", "gemini-2.5-flash")
        )
        normalized["model_override"] = str(normalized.get("model_override", "") or "")
        row["settings"] = {key: normalized.get(key) for key in ROW_SETTING_KEYS}
        row["engine"] = self.LabelFromModel(row["settings"].get("model_name", "gemini-2.5-flash"))
        return row["settings"]

    def FormatRowSettingsSummary(self, settings):
        fmt = str(settings.get("format_type", "html")).upper()
        profile_key = settings.get("doc_profile", "standard")
        profile_label = PROFILE_KEY_TO_LABEL.get(profile_key, "Standard")
        punct = "strict punctuation" if not settings.get("modernize_punctuation", False) else "modern punctuation"
        translate_mode = str(settings.get("translate_mode", "none"))
        trans = "no translation" if translate_mode == "none" else ("translate-only" if translate_mode == "english_only" else "original+translation")
        units = "units untouched" if not settings.get("unit_conversion", False) else "units converted"
        page_refs = "printed page refs on" if settings.get("preserve_original_page_numbers", False) else "printed page refs off"
        parts = [fmt, profile_label, punct, trans, units, page_refs]
        pdf_scope = normalize_pdf_page_scope_text(settings.get("pdf_page_scope", ""))
        if pdf_scope:
            parts.append(f"PDF pages {pdf_scope}")
        return " | ".join(parts)

    def GetTargetQueueIndicesForSettingChange(self):
        return shared_get_target_queue_indices_for_setting_change(
            self.queue,
            self.GetSelectedIndices(),
        )

    def SyncRuntimeControlsFromSelection(self):
        selected = self.GetSelectedIndices()
        if not selected:
            return
        idx = selected[0]
        if not (0 <= idx < len(self.queue)):
            return
        settings = self.NormalizeRowSettings(self.queue[idx])
        tmap = {'none': 0, 'both': 1, 'english_only': 2}
        fmt_map = {'html': 0, 'txt': 1, 'docx': 2, 'md': 3, 'pdf': 4, 'json': 5, 'csv': 6, 'epub': 7}
        profile_key = settings.get('doc_profile', 'standard')
        profile_index = next((i for i, (k, _) in enumerate(PROFILE_CHOICES) if k == profile_key), 0)
        self.format_choice.SetSelection(fmt_map.get(settings.get('format_type', 'html'), 0))
        self.choice_profile.SetSelection(profile_index)
        self.run_translate_choice.SetSelection(tmap.get(settings.get('translate_mode', 'none'), 0))
        self.run_punct_choice.SetSelection(1 if settings.get('modernize_punctuation', False) else 0)
        self.run_units_choice.SetSelection(1 if settings.get('unit_conversion', False) else 0)
        self.run_abbrev_choice.SetSelection(1 if settings.get('abbrev_expansion', False) else 0)
        self.run_images_choice.SetSelection(0 if settings.get('image_descriptions', True) else 1)
        self.run_page_numbers_choice.SetSelection(1 if settings.get('preserve_original_page_numbers', False) else 0)
        self.run_merge_choice.SetSelection(1 if settings.get('merge_files', False) else 0)
        self.run_largeprint_choice.SetSelection(1 if settings.get('large_print', False) else 0)
        self.SetPdfPageScopeSelection(settings.get('pdf_page_scope', '') or '')
        self.UpdateModelGuide()
        self.UpdateProfileGuide()
        self.UpdateLargePrintVisibility()
        self.UpdatePreflightSummary()

    def UpdateLargePrintVisibility(self):
        selected_fmt = self.format_choice.GetStringSelection().strip().lower() if self.format_choice else "html"
        show_pdf_controls = selected_fmt == "pdf"
        self.lbl_run_largeprint.Show(show_pdf_controls)
        self.run_largeprint_choice.Show(show_pdf_controls)
        parent = self.run_largeprint_choice.GetParent()
        if parent:
            parent.Layout()

    def QueueDisplayString(self, row):
        name = os.path.basename(row.get("path", "")) or "Untitled item"
        status = self.GetQueueDisplayStatus(row)
        settings = self.NormalizeRowSettings(row)
        return f"{name} | Status: {status} | {self.FormatRowSettingsSummary(settings)}"

    def GetQueueDisplayStatus(self, row):
        if hasattr(self, "queue_panel") and self.queue_panel is not None:
            return self.queue_panel.GetQueueDisplayStatus(row)
        return get_queue_display_status(row)

    def GetQueueCurrentRowIndex(self):
        return self.queue_panel.GetQueueCurrentRowIndex()

    def BuildQueueCurrentRowAnnouncement(self):
        if hasattr(self, "queue_panel") and self.queue_panel is not None:
            return self.queue_panel.BuildQueueCurrentRowAnnouncement()
        return build_queue_current_row_announcement(self)

    def EnsureQueueTableLanding(self, focus=False, select_row=True):
        if hasattr(self, "queue_panel") and self.queue_panel is not None:
            return self.queue_panel.EnsureQueueTableLanding(focus=focus, select_row=select_row)
        return ensure_queue_table_landing(self, focus=focus, select_row=select_row)

    def _queue_has_real_items(self):
        return len(self.queue) > 0

    def _has_queue_selection(self):
        return len(self.GetSelectedIndices()) > 0

    def _clear_queue_selection(self):
        if not hasattr(self, "queue_panel"):
            return
        self.queue_panel.ClearQueueSelection()

    def _select_queue_rows(self, rows):
        if not hasattr(self, "queue_panel"):
            return
        self.queue_panel.SelectQueueRows(rows)

    def RefreshQueueRows(self, indices):
        if not hasattr(self, "queue_panel"):
            return
        self.queue_panel.RefreshQueueRows(indices)

    def ApplySettingsToRows(self, indices, settings, reason_label):
        if not indices:
            self.Log(f"No queued/paused rows available to apply {reason_label}.", engine_event=True)
            return
        shared_apply_settings_to_rows(
            self.queue,
            indices,
            settings,
            row_setting_keys=ROW_SETTING_KEYS,
            label_from_model_fn=self.LabelFromModel,
        )
        self.RefreshQueueRows(indices)
        self.SaveActiveSession()
        selection_scope = "selected rows" if self.GetSelectedIndices() else "all queued/paused rows"
        self.Log(f"Applied {reason_label} to {len(indices)} {selection_scope}.", engine_event=True)

    def GetProfileGuideText(self, profile_key):
        return profile_tooltip_text(profile_key).splitlines()[0]

    def UpdateProfileGuide(self):
        profile_key = PROFILE_CHOICES[self.choice_profile.GetSelection()][0]
        profile_label = PROFILE_CHOICES[self.choice_profile.GetSelection()][1]
        preferred_model_name = self.GetPreferredModelName(profile_key)
        model_name = self.ResolveModelNameForAvailableKeys(preferred_model_name)
        speed_warning = get_processing_speed_warning(profile_key, model_name)
        tooltip = profile_tooltip_text(profile_key)
        if str(self.cfg.get("model_override", "") or "").strip():
            tooltip = f"{tooltip}\n\nPreference override: {self.LabelFromModel(preferred_model_name)}."
        if model_name != preferred_model_name and any(self.keys.values()):
            preferred_label = self.LabelFromModel(preferred_model_name)
            tooltip = (
                f"{tooltip}\n\nEngine fallback: {self.LabelFromModel(model_name)} will be used until a key is added for {preferred_label}."
            )
        if speed_warning:
            tooltip = f"{tooltip}\n\n{speed_warning}"
        self.choice_profile.SetToolTip(tooltip)
        self.choice_profile.SetName(f"Document Preset Picker ({profile_label})")
        if hasattr(self, "btn_start"):
            start_tooltip = "Start reading everything that is currently queued."
            if speed_warning:
                start_tooltip = f"{start_tooltip} {speed_warning}"
            self.btn_start.SetToolTip(start_tooltip)

    def OnProfileChoiceChanged(self, event):
        profile_key = PROFILE_CHOICES[self.choice_profile.GetSelection()][0]
        self.ApplyProfilePresetToControls(profile_key)
        self.UpdateModelGuide()
        self.UpdateProfileGuide()
        self.UpdatePreflightSummary()
        if not self.is_running and self.queue:
            settings = self.BuildSettingsFromControls()
            target_indices = self.GetTargetQueueIndicesForSettingChange()
            self.ApplySettingsToRows(target_indices, settings, "current profile settings")
        event.Skip()

    def OnFormatChoiceChanged(self, event):
        self.UpdateLargePrintVisibility()
        self.UpdatePreflightSummary()
        if not self.is_running and self.queue:
            settings = self.BuildSettingsFromControls()
            target_indices = self.GetTargetQueueIndicesForSettingChange()
            self.ApplySettingsToRows(target_indices, settings, "current format settings")
        event.Skip()

    def OnRunOptionChoiceChanged(self, event):
        self.UpdatePreflightSummary()
        if not self.is_running and self.queue:
            settings = self.BuildSettingsFromControls()
            target_indices = self.GetTargetQueueIndicesForSettingChange()
            self.ApplySettingsToRows(target_indices, settings, "run-time option settings")
        event.Skip()

    def OnPdfPageScopeChanged(self, event):
        if self.choice_pdf_pages.GetStringSelection() == "Custom...":
            custom_scope = self.PromptForCustomPdfPageScope()
            if custom_scope is None:
                self.SetPdfPageScopeSelection("")
            else:
                self.SetPdfPageScopeSelection(custom_scope)
        self.UpdatePreflightSummary()
        if not self.is_running and self.queue:
            settings = self.BuildSettingsFromControls()
            target_indices = self.GetTargetQueueIndicesForSettingChange()
            self.ApplySettingsToRows(target_indices, settings, "PDF page scope settings")
        event.Skip()

    def OnApplySettingsToQueue(self, event):
        if self.is_running:
            return
        settings = self.BuildSettingsFromControls()
        target_indices = self.GetTargetQueueIndicesForSettingChange()
        self.ApplySettingsToRows(target_indices, settings, "current reading settings")
        event.Skip()

    def GetSelectedIndices(self):
        if not hasattr(self, "queue_panel"):
            return []
        return self.queue_panel.GetSelectedIndices()

    def UpdateQueueButtons(self):
        has_items = len(self.queue) > 0
        has_sel = self._has_queue_selection()
        can_edit = not self.is_running
        self.format_choice.Enable(can_edit)
        self.run_translate_choice.Enable(can_edit)
        self.run_punct_choice.Enable(can_edit)
        self.run_units_choice.Enable(can_edit)
        self.run_abbrev_choice.Enable(can_edit)
        self.run_images_choice.Enable(can_edit)
        self.run_merge_choice.Enable(can_edit)
        self.run_largeprint_choice.Enable(can_edit)
        self.choice_pdf_pages.Enable(can_edit)
        self.btn_add.Enable(can_edit)
        self.btn_add_folder.Enable(can_edit)
        self.btn_remove.Enable(can_edit and has_sel)
        self.btn_task_actions.Enable(has_sel)
        self.btn_select_all.Enable(has_items)
        self.btn_deselect_all.Enable(has_sel)
        self.btn_clear.Enable(can_edit and has_items)
        self.btn_apply_settings.Enable(can_edit and has_items)
        if hasattr(self, "btn_run_preflight") and self.btn_run_preflight is not None:
            self.btn_run_preflight.Enable(can_edit and has_items)
        self.btn_start.Enable(can_edit and has_items)
        self.btn_schedule.Enable(can_edit and has_items)
        self.choice_recursive.Enable(can_edit)
        self.choice_profile.Enable(can_edit)
        self.chk_preserve_structure.Enable(can_edit)
        self.chk_delete_originals.Enable(can_edit)
        self.dest_choice.Enable(can_edit)
        self.txt_dest.Enable(can_edit and self.dest_choice.GetSelection() == 1)
        self.btn_dest.Enable(can_edit and self.dest_choice.GetSelection() == 1)

    def _estimate_path_work_units(self, path, settings=None):
        return shared_estimate_path_work_units(
            path,
            settings=settings,
            pdf_reader_factory=PdfReader,
            normalize_pdf_page_scope_text_fn=normalize_pdf_page_scope_text,
            parse_pdf_page_scope_spec_fn=parse_pdf_page_scope_spec,
        )

    def RefreshQueueWorkUnitEstimates(self):
        shared_refresh_queue_work_unit_estimates(
            self.queue,
            normalize_row_settings_fn=self.NormalizeRowSettings,
            estimate_path_work_units_fn=self._estimate_path_work_units,
        )

    def GetRunUnitTotals(self):
        return shared_get_run_unit_totals(
            self.queue,
            current_processing_index=self.current_processing_index,
            current_file_page_done=self.current_file_page_done,
            terminal_statuses=SESSION_TERMINAL_STATUSES,
        )

    def ShouldLogPageProgress(self, done_pages, total_pages):
        return shared_should_log_page_progress(done_pages, total_pages)

    def ShouldStatusEchoLog(self, msg, engine_event=False):
        return shared_should_status_echo_log(msg, engine_event=engine_event)

    def _format_engine_log_message(self, msg, now=None):
        now = time.time() if now is None else now
        start_ts = getattr(self, "current_run_started_ts", 0.0) or now
        elapsed = max(0, int(now - start_ts))
        clock = time.strftime("%H:%M:%S", time.localtime(now))
        hours = elapsed // 3600
        minutes = (elapsed % 3600) // 60
        seconds = elapsed % 60
        return f"[{clock} | +{hours:02d}:{minutes:02d}:{seconds:02d}] {msg}"

    def _maybe_log_progress_heartbeat(self, now=None):
        now = time.time() if now is None else now
        if not getattr(self, "is_running", False) or getattr(self, "is_paused", False):
            return False
        current_idx = getattr(self, "current_processing_index", -1)
        if not (0 <= current_idx < len(getattr(self, "queue", []))):
            return False
        current_row = self.queue[current_idx]
        if str(current_row.get("status", "")) != "Processing":
            return False
        last_activity = getattr(self, "_last_engine_activity_ts", 0.0)
        last_heartbeat = getattr(self, "_last_progress_heartbeat_ts", 0.0)
        if now - last_activity < 30.0 or now - last_heartbeat < 30.0:
            return False
        total_pages = max(0, int(getattr(self, "current_file_page_total", 0) or 0))
        done_pages = max(0, int(getattr(self, "current_file_page_done", 0) or 0))
        recovered_units = max(0, int(getattr(self, "current_file_resume_recovered_units", 0) or 0))
        remaining_units = max(0, int(getattr(self, "current_file_resume_remaining_units", 0) or 0))
        file_name = getattr(self, "current_file_name", "") or os.path.basename(str(current_row.get("path", "")))
        if total_pages > 0:
            unit_label = getattr(self, "current_file_unit_label", "items") or "items"
            if recovered_units > 0 and remaining_units > 0:
                resume_done = max(0, done_pages - recovered_units)
                self.Log(
                    f"[Progress] Total progress for {file_name}: {done_pages}/{total_pages} {unit_label} complete "
                    f"(resume pass: {resume_done}/{remaining_units} remaining {unit_label} complete).",
                    engine_event=True,
                )
            else:
                self.Log(f"[Progress] Still working on {file_name}: {done_pages}/{total_pages} {unit_label} processed.", engine_event=True)
        else:
            self.Log(f"[Progress] Still working on {file_name}.", engine_event=True)
        self._last_progress_heartbeat_ts = now
        return True

    def _update_pdf_progress_state(self, done_pages, total_pages):
        self.current_file_page_done = done_pages
        self.current_file_page_total = max(1, total_pages)
        self.total_pages_processed += 1
        return self.total_pages_processed

    def _persist_resume_state(self, done_pages, total_pages, current_source_page=None):
        progress_path = getattr(self, "current_file_resume_state_path", "")
        if not progress_path:
            return
        payload = {
            "completed_units": max(0, int(done_pages or 0)),
            "total_units": max(0, int(total_pages or 0)),
            "completed_pages": max(0, int(done_pages or 0)),
            "total_pages": max(0, int(total_pages or 0)),
            "current_source_unit": int(current_source_page) if current_source_page is not None else None,
            "current_source_page": int(current_source_page) if current_source_page is not None else None,
            "updated_at": time.time(),
        }
        try:
            write_progress_state(progress_path, payload)
        except Exception as ex:
            self.Log(f"Warning: could not persist PDF resume state ({ex})", engine_event=True)

    def _persist_merge_resume_state(self, progress_temp_path, completed_job_paths, total_units):
        if not progress_temp_path:
            return
        payload = {
            "completed_job_paths": list(completed_job_paths),
            "completed_units": len(completed_job_paths),
            "total_units": max(0, int(total_units or 0)),
            "updated_at": time.time(),
        }
        try:
            write_progress_state(progress_temp_path, payload)
        except Exception as ex:
            self.Log(f"Warning: could not persist merge resume state ({ex})", engine_event=True)

    def BuildProgressSummary(self):
        return shared_build_progress_summary(
            self.queue,
            current_processing_index=getattr(self, "current_processing_index", -1),
            current_file_ordinal=getattr(self, "current_file_ordinal", 0),
            current_file_page_total=getattr(self, "current_file_page_total", 0),
            current_file_page_done=getattr(self, "current_file_page_done", 0),
            terminal_statuses=SESSION_TERMINAL_STATUSES,
        )

    def UpdateProgressIndicators(self, force_status=False):
        total_units, completed_units = self.GetRunUnitTotals()
        gauge_value = int((completed_units / total_units) * 100) if total_units > 0 else 0
        self.progress_gauge.SetValue(gauge_value)
        summary = self.BuildProgressSummary()
        self.progress_summary.SetValue(summary)
        # Keep status updates informative but not chatty so screen readers are not constantly interrupted.
        now = time.time()
        last_ts = getattr(self, "_last_progress_status_ts", 0.0)
        if force_status or (now - last_ts >= 2.0):
            self.SafeSetStatusText(summary[:180])
            self._last_progress_status_ts = now
        self._maybe_log_progress_heartbeat(now=now)

    def RefreshQueue(self):
        if hasattr(self, "queue_panel") and self.queue_panel is not None:
            self.queue_panel.RefreshQueue()
            if hasattr(self, "UpdatePreflightSummary"):
                self.UpdatePreflightSummary()
            return
        file_list = getattr(self, "file_list", None)
        if file_list is None:
            return
        if hasattr(file_list, "Clear"):
            file_list.Clear()
        elif hasattr(file_list, "DeleteAllItems"):
            file_list.DeleteAllItems()
        if not getattr(self, "queue", None):
            if hasattr(file_list, "Append"):
                file_list.Append(QUEUE_EMPTY_PLACEHOLDER)
            elif hasattr(file_list, "AppendItem"):
                file_list.AppendItem([QUEUE_EMPTY_PLACEHOLDER])
            if hasattr(file_list, "UnselectAll"):
                file_list.UnselectAll()
        else:
            for row in self.queue:
                row_text = self.QueueDisplayString(row)
                if hasattr(file_list, "Append"):
                    file_list.Append(row_text)
                elif hasattr(file_list, "AppendItem"):
                    file_list.AppendItem([row_text])
        self.UpdateQueueButtons()
        self.UpdateProgressIndicators()
        if hasattr(self, "UpdatePreflightSummary"):
            self.UpdatePreflightSummary()

    def Log(self, msg, engine_event=False):
        rendered = msg
        if engine_event:
            now = time.time()
            if not getattr(self, "current_run_started_ts", 0.0):
                self.current_run_started_ts = now
            self._last_engine_activity_ts = now
            rendered = self._format_engine_log_message(msg, now=now)
            self.processing_log_lines.append(rendered)
            wx.CallAfter(self.lt.AppendText, f"{rendered}\n")
        if self.ShouldStatusEchoLog(msg, engine_event=engine_event):
            wx.CallAfter(self.SafeSetStatusText, msg[:180])
        print(rendered)

    def GetDefaultLogDirectory(self):
        return shared_resolve_default_log_directory(self.queue, self.cfg, SCRIPT_DIR)

    def AutoSaveProcessingLog(self):
        if not self.processing_log_lines:
            return None
        if not bool(self.cfg.get('auto_save_processing_log', False)):
            return None
        log_dir = self.GetDefaultLogDirectory()
        return shared_write_processing_log(log_dir, self.build_stamp, self.processing_log_lines)

    def OnSaveLog(self, e):
        if not self.processing_log_lines:
            wx.MessageBox("No processing log entries available yet.", "Save Processing Log", wx.OK | wx.ICON_INFORMATION)
            return
        default_dir = self.GetDefaultLogDirectory()
        stamp = time.strftime("%Y%m%d_%H%M%S")
        default_name = f"chronicle_processing_log_{stamp}.txt"
        dlg = wx.FileDialog(
            self,
            "Save processing log",
            defaultDir=default_dir,
            defaultFile=default_name,
            wildcard="Text files (*.txt)|*.txt|Log files (*.log)|*.log|All files (*.*)|*.*",
            style=wx.FD_SAVE | wx.FD_OVERWRITE_PROMPT,
        )
        if dlg.ShowModal() == wx.ID_OK:
            path = dlg.GetPath()
            try:
                with open(path, 'w', encoding='utf-8') as fh:
                    lines = build_log_header(self.build_stamp) + self.processing_log_lines
                    fh.write("\n".join(lines).rstrip() + "\n")
                self.Log(f"Processing log saved: {path}")
            except Exception as ex:
                wx.MessageBox(f"Could not save log file.\n\nDetails: {ex}", "Save Log Error", wx.OK | wx.ICON_ERROR)
        dlg.Destroy()

    def SetQueueStatus(self, index, status):
        if 0 <= index < len(self.queue):
            self.queue[index]['status'] = status
            wx.CallAfter(self.RefreshQueue)
            wx.CallAfter(self.UpdateProgressIndicators)
            self.SaveActiveSession()

    def SetRunningState(self, running):
        for key, value in shared_build_running_state_update(running).items():
            setattr(self, key, value)
        self.UpdateQueueButtons()
        self.UpdateProgressIndicators(force_status=True)

    def OnQueueSelectionChanged(self, e):
        self.UpdateQueueButtons()
        if not self.is_running:
            self.SyncRuntimeControlsFromSelection()
            self.UpdatePreflightSummary()
        e.Skip()

    def OnQueueKeyDown(self, e):
        code = e.GetKeyCode()
        if (e.CmdDown() or e.ControlDown()) and (code == ord('A') or code == ord('a')):
            self.OnSelectAllQueue(None)
            return
        if code == wx.WXK_ESCAPE:
            self.OnDeselectAllQueue(None)
            return
        if code == wx.WXK_SPACE:
            self.ToggleCurrentQueueRowSelection()
            return
        if code == wx.WXK_DELETE:
            self.OnDeleteTask(None)
            return
        if e.ShiftDown() and code == wx.WXK_F10:
            self.ShowQueueContextMenu()
            return
        if hasattr(wx, "WXK_MENU") and code == wx.WXK_MENU:
            self.ShowQueueContextMenu()
            return
        if hasattr(wx, "WXK_WINDOWS_MENU") and code == wx.WXK_WINDOWS_MENU:
            self.ShowQueueContextMenu()
            return
        e.Skip()

    def ShowQueueContextMenu(self):
        menu = wx.Menu()
        id_stop = wx.NewIdRef()
        id_pause = wx.NewIdRef()
        id_resume = wx.NewIdRef()
        id_delete = wx.NewIdRef()
        id_open = wx.NewIdRef()
        menu.Append(id_stop, "Stop Task")
        menu.Append(id_pause, "Pause Task")
        menu.Append(id_resume, "Resume Task")
        menu.AppendSeparator()
        menu.Append(id_delete, "Delete Task")
        menu.Append(id_open, "Open Containing Folder")

        selected = self.GetSelectedIndices()
        has_sel = len(selected) > 0
        can_stop = False
        can_pause = False
        can_resume = False
        for idx in selected:
            status = str(self.queue[idx].get("status", "Queued"))
            if status in {"Queued", "Processing", "Paused"}:
                can_stop = True
            if status in {"Queued", "Processing"}:
                can_pause = True
            if status == "Paused":
                can_resume = True
        menu.Enable(id_stop, has_sel and can_stop)
        menu.Enable(id_pause, has_sel and can_pause)
        menu.Enable(id_resume, has_sel and can_resume)
        menu.Enable(id_delete, has_sel and not self.is_running)
        menu.Enable(id_open, has_sel)

        self.Bind(wx.EVT_MENU, self.OnStopTask, id=id_stop)
        self.Bind(wx.EVT_MENU, self.OnPauseTask, id=id_pause)
        self.Bind(wx.EVT_MENU, self.OnResumeTask, id=id_resume)
        self.Bind(wx.EVT_MENU, self.OnDeleteTask, id=id_delete)
        self.Bind(wx.EVT_MENU, self.OnOpenContainingFolder, id=id_open)
        self.PopupMenu(menu)
        menu.Destroy()

    def OnTaskActionsButton(self, e):
        self.ShowQueueContextMenu()

    def _unselect_row(self, row):
        if 0 <= row < len(self.queue):
            self.file_list.Deselect(row)

    def OnSelectAllQueue(self, e):
        if self.queue:
            self._select_queue_rows(range(len(self.queue)))
            self.SetQueueCurrentRow(0)
        else:
            self._clear_queue_selection()
        self.UpdateQueueButtons()
        if e:
            e.Skip()

    def OnDeselectAllQueue(self, e):
        self._clear_queue_selection()
        self.UpdateQueueButtons()
        if e:
            e.Skip()

    def ToggleCurrentQueueRowSelection(self):
        selections = self.GetSelectedIndices()
        row = self.file_list.GetSelection()
        if row == wx.NOT_FOUND and selections:
            row = selections[0]
        if row == wx.NOT_FOUND or row < 0:
            return
        if row in selections:
            self._unselect_row(row)
        else:
            self.file_list.SetSelection(row)
        self.UpdateQueueButtons()

    def OnQueueItemContextMenu(self, e):
        if self.queue and not self.GetSelectedIndices():
            self.file_list.SetSelection(0)
        self.ShowQueueContextMenu()

    def OnStopTask(self, e):
        selected = self.GetSelectedIndices()
        if not selected:
            return
        stop_result = shared_stop_selected_tasks(self.queue, selected)
        stopped = stop_result["stopped"]
        includes_processing = stop_result["includes_processing"]
        if includes_processing and self.is_running:
            self.stop_requested = True
            self.is_paused = False
            self.Log("Stop requested for active task. Chronicle will stop safely at the next processing boundary.", engine_event=True)
        if stopped:
            self.RefreshQueue()
            self.SaveActiveSession()
            self.Log(f"Stopped {stopped} queued/paused task(s).")

    def OnDeleteTask(self, e):
        if self.is_running:
            return
        self.OnRemoveSelected(e)

    def OnPauseTask(self, e):
        selected = self.GetSelectedIndices()
        if not selected:
            return
        paused = shared_pause_selected_tasks(self.queue, selected)
        if self.is_running:
            self.is_paused = True
        if paused:
            self.RefreshQueue()
            self.SaveActiveSession()
            self.Log(f"Paused {paused} selected task(s).")

    def OnResumeTask(self, e):
        selected = self.GetSelectedIndices()
        if not selected:
            return
        resumed = shared_resume_selected_tasks(self.queue, selected)
        if resumed:
            self.is_paused = False
            self.RefreshQueue()
            self.SaveActiveSession()
            self.Log(f"Resumed {resumed} selected task(s).")

    def OnOpenContainingFolder(self, e):
        selected = self.GetSelectedIndices()
        if not selected:
            return
        idx = selected[0]
        if not (0 <= idx < len(self.queue)):
            return
        folder = os.path.dirname(self.queue[idx]['path'])
        if not os.path.isdir(folder):
            return
        try:
            if platform.system() == "Darwin":
                subprocess.Popen(["open", folder])
            elif platform.system() == "Windows":
                os.startfile(folder)
            else:
                subprocess.Popen(["xdg-open", folder])
            self.Log(f"Opened folder: {folder}")
        except Exception as ex:
            self.Log(f"Unable to open folder: {ex}")

    def OnFileActivated(self, e):
        selected = self.GetSelectedIndices()
        idx = selected[0] if selected else wx.NOT_FOUND
        if idx < 0 or idx >= len(self.queue):
            return
        folder = os.path.dirname(self.queue[idx]['path'])
        if not os.path.isdir(folder):
            return
        try:
            if platform.system() == "Darwin":
                subprocess.Popen(["open", folder])
            elif platform.system() == "Windows":
                os.startfile(folder)
            else:
                subprocess.Popen(["xdg-open", folder])
            self.Log(f"Opened folder: {folder}")
        except Exception as ex:
            self.Log(f"Unable to open folder: {ex}")

    def OnDestModeChanged(self, e=None):
        is_custom = self.dest_choice.GetSelection() == 1
        self.txt_dest.Enable(is_custom and not self.is_running)
        self.btn_dest.Enable(is_custom and not self.is_running)
        if e:
            e.Skip()

    def GetBrowseStartDir(self):
        candidate = str(self.cfg.get('last_browse_dir', '')).strip()
        if candidate and os.path.isdir(candidate):
            return candidate
        custom_dest = str(self.cfg.get('custom_dest', '')).strip()
        if custom_dest and os.path.isdir(custom_dest):
            return custom_dest
        docs_dir = os.path.expanduser('~/Documents')
        if os.path.isdir(docs_dir):
            return docs_dir
        if os.path.isdir(SCRIPT_DIR):
            return SCRIPT_DIR
        return os.path.expanduser('~')

    def RememberBrowsePath(self, path):
        if not path:
            return
        candidate = path if os.path.isdir(path) else os.path.dirname(path)
        if candidate and os.path.isdir(candidate):
            self.cfg['last_browse_dir'] = candidate
            save_json(CONFIG_FILE, self.cfg)

    def _run_osascript_dialog(self, script_lines):
        if platform.system() != "Darwin":
            raise RuntimeError("AppleScript dialog fallback is only available on macOS.")
        cmd = ["osascript"]
        for line in script_lines:
            cmd.extend(["-e", line])
        proc = subprocess.run(cmd, capture_output=True, text=True)
        if proc.returncode != 0:
            message = (proc.stderr or proc.stdout or "Unknown AppleScript error").strip()
            if any(token in message.lower() for token in ("user canceled", "cancelled", "(-128)")):
                return None
            raise RuntimeError(message)
        output = proc.stdout.strip()
        return output or None

    def _mac_choose_files(self):
        script_lines = [
            'set chosenFiles to choose file with prompt "Select files to add" multiple selections allowed true',
            'set outputLines to {}',
            'repeat with currentFile in chosenFiles',
            'copy POSIX path of currentFile to end of outputLines',
            'end repeat',
            'set text item delimiters of AppleScript to linefeed',
            'return outputLines as text',
        ]
        output = self._run_osascript_dialog(script_lines)
        if output is None:
            return None
        return [line for line in output.splitlines() if line.strip()]

    def _mac_choose_folder(self, message):
        safe_message = str(message).replace('"', '\"')
        script_lines = [
            f'set chosenFolder to choose folder with prompt "{safe_message}"',
            'return POSIX path of chosenFolder',
        ]
        return self._run_osascript_dialog(script_lines)


    def _build_file_dialog(self, message, wildcard, style):
        start_dir = self.GetBrowseStartDir()
        kwargs = {
            'message': message,
            'wildcard': wildcard,
            'style': style,
        }
        if start_dir and os.path.isdir(start_dir):
            kwargs['defaultDir'] = start_dir
        try:
            return wx.FileDialog(self, **kwargs)
        except TypeError:
            kwargs.pop('defaultDir', None)
            return wx.FileDialog(self, **kwargs)

    def _build_dir_dialog(self, message):
        start_dir = self.GetBrowseStartDir()
        kwargs = {
            'message': message,
            'style': wx.DD_DEFAULT_STYLE | wx.DD_DIR_MUST_EXIST,
        }
        if start_dir and os.path.isdir(start_dir):
            kwargs['defaultPath'] = start_dir
        try:
            return wx.DirDialog(self, **kwargs)
        except TypeError:
            kwargs.pop('defaultPath', None)
            return wx.DirDialog(self, **kwargs)

    def OnChooseDest(self, e):
        try:
            self.Raise()
        except Exception:
            pass
        dlg = self._build_dir_dialog("Select output folder")
        if dlg.ShowModal() == wx.ID_OK:
            chosen = dlg.GetPath()
            self.txt_dest.SetValue(chosen)
            self.RememberBrowsePath(chosen)
        dlg.Destroy()

    def OnSafetyOptionsChanged(self, e=None):
        self.cfg['preserve_source_structure'] = bool(self.chk_preserve_structure.GetValue())
        self.cfg['delete_source_on_success'] = bool(self.chk_delete_originals.GetValue())
        save_json(CONFIG_FILE, self.cfg)
        self.SaveActiveSession()
        if e:
            e.Skip()

    def AddPathEntries(self, paths, source_root=None):
        settings = self.BuildSettingsFromControls()
        engine = self.LabelFromModel(settings.get("model_name", "gemini-2.5-flash"))
        return shared_add_path_entries(
            self.queue,
            paths,
            settings=settings,
            engine_label=engine,
            row_setting_keys=ROW_SETTING_KEYS,
            source_root=source_root,
        )

    def SelectQueueRowsByPaths(self, paths):
        selected_rows = shared_find_queue_rows_by_paths(self.queue, paths)
        if selected_rows:
            self._select_queue_rows(selected_rows)
            self.SetQueueCurrentRow(selected_rows[0])
        self.UpdateQueueButtons()

    def SetQueueCurrentRow(self, row, focus=False):
        if not (0 <= row < len(self.queue)):
            return
        if row not in self.GetSelectedIndices():
            self.file_list.SetSelection(row)
        try:
            self.file_list.SetFirstItem(row)
        except Exception:
            pass
        if focus:
            self.file_list.SetFocus()

    def GetPageSequenceNumber(self, path):
        return shared_get_page_sequence_number(path)

    def GetOrderedJobsForProcessing(self):
        jobs, sequence_lock = shared_get_ordered_jobs_for_processing(
            self.queue,
            merge_files=bool(self.cfg.get('merge_files', False)),
            page_sequence_fn=self.GetPageSequenceNumber,
        )
        if sequence_lock:
            self.Log("[Merge] Page-sequence lock enabled for merge mode.", engine_event=True)
        return jobs

    def ResolveMergeOutputPath(self, jobs, fmt, custom_dest, dest_mode):
        return shared_resolve_merge_output_path(
            jobs,
            fmt,
            custom_dest=custom_dest,
            dest_mode=dest_mode,
            script_dir=SCRIPT_DIR,
            collision_mode=self.cfg.get('collision_mode', 'skip'),
        )

    def OnAddFiles(self, e):
        wildcard = (
            'Supported Files|*.pdf;*.docx;*.txt;*.md;*.rtf;*.csv;*.js;*.jpg;*.jpeg;*.png;'
            '*.bmp;*.tiff;*.tif;*.webp;*.epub;*.pptx;*.ppt;*.xlsx;*.xls'
        )
        dlg = wx.FileDialog(self, 'Select files to add', wildcard=wildcard, style=wx.FD_OPEN | wx.FD_FILE_MUST_EXIST | wx.FD_MULTIPLE)
        if dlg.ShowModal() == wx.ID_OK:
            selected_paths = dlg.GetPaths()
            if selected_paths:
                self.RememberBrowsePath(selected_paths[0])
                added_paths = self.AddPathEntries(selected_paths, source_root=None)
                self.RefreshQueue()
                self.SelectQueueRowsByPaths(added_paths)
                self.SaveActiveSession()
                self.Log(f"Added {len(added_paths)} file(s) to the file list.")
        dlg.Destroy()
        self.file_list.SetFocus()

    def OnAddFolder(self, e):
        dlg = wx.DirDialog(self, "Select folder to scan", style=wx.DD_DEFAULT_STYLE | wx.DD_DIR_MUST_EXIST)
        if dlg.ShowModal() != wx.ID_OK:
            dlg.Destroy()
            self.file_list.SetFocus()
            return
        folder = dlg.GetPath()
        dlg.Destroy()
        self.RememberBrowsePath(folder)

        paths = []
        try:
            if self.choice_recursive.GetSelection() == 1:
                for root_dir, _, files in os.walk(folder):
                    for name in files:
                        if os.path.splitext(name)[1].lower() in SUPPORTED_EXTENSIONS and not name.startswith('.'):
                            paths.append(os.path.join(root_dir, name))
            else:
                for name in os.listdir(folder):
                    full = os.path.join(folder, name)
                    if not os.path.isfile(full):
                        continue
                    if os.path.splitext(name)[1].lower() in SUPPORTED_EXTENSIONS and not name.startswith('.'):
                        paths.append(full)
        except PermissionError as ex:
            wx.MessageBox(
                "Chronicle does not currently have permission to read this folder.\n\n"
                "Use File > Check Mac Folder Access to grant access, then try again.\n\n"
                f"Details: {ex}",
                "Folder Permission Denied",
                wx.OK | wx.ICON_WARNING,
            )
            self.file_list.SetFocus()
            return

        added_paths = self.AddPathEntries(sorted(paths), source_root=folder)
        self.RefreshQueue()
        self.SelectQueueRowsByPaths(added_paths)
        self.SaveActiveSession()
        self.Log(f"Scanned folder and added {len(added_paths)} file(s): {folder}")
        self.file_list.SetFocus()

    def OnRemoveSelected(self, e):
        selected = self.GetSelectedIndices()
        if not selected:
            return
        for idx in reversed(selected):
            del self.queue[idx]
        self.RefreshQueue()
        self.SaveActiveSession()
        if self.queue:
            next_idx = min(selected[0], len(self.queue) - 1)
            self.SetQueueCurrentRow(next_idx)
        self.Log(f"Removed {len(selected)} selected item(s).")
        self.file_list.SetFocus()

    def OnClearQueue(self, e):
        count = len(self.queue)
        self.queue = []
        self.RefreshQueue()
        self.SaveActiveSession()
        self.Log(f"Cleared file list ({count} item(s) removed).")
        self.file_list.SetFocus()

    def OnApiKeys(self, e):
        d = ApiKeyDialog(self, initial_keys=self.keys, save_keys=lambda keys: save_json(KEYS_FILE, keys))
        result = d.ShowModal()
        if result == wx.ID_OK:
            self.keys = dict(d.keys)
        d.Destroy()
        self.clients = {}
        if not self.is_running:
            self.cfg["model_name"] = self.ResolveModelNameForAvailableKeys(
                self.GetPreferredModelName(self.cfg.get("doc_profile", "standard"))
            )
            for row in self.queue:
                self.NormalizeRowSettings(row)
            self._sync_controls_from_cfg()
            self.RefreshQueue()
            self.SaveActiveSession()

    def OnDiscoverScanners(self, e):
        self.Log("Searching for connected input devices...", engine_event=True)
        scanners, err = discover_connected_flatbed_scanners()
        if err:
            self.Log(err, engine_event=True)
        if not scanners:
            msg = "No connected input devices were detected."
            if err:
                msg += f"\n\nDetails: {err}"
            wx.MessageBox(msg, "Device Search", wx.OK | wx.ICON_INFORMATION)
            return

        lines = []
        for idx, scanner in enumerate(scanners, start=1):
            name = scanner.get("name", "Unknown device")
            manufacturer = scanner.get("manufacturer", "").strip()
            source = scanner.get("source", "").strip()
            extras = []
            if manufacturer:
                extras.append(manufacturer)
            if source:
                extras.append(source)
            suffix = f" ({', '.join(extras)})" if extras else ""
            lines.append(f"{idx}. {name}{suffix}")
        summary = "\n".join(lines)
        self.Log(f"Device search complete: {len(scanners)} device(s) found.", engine_event=True)
        wx.MessageBox(summary, "Connected Devices", wx.OK | wx.ICON_INFORMATION)

    def OnScanViaNaps2(self, e):
        if self.is_running:
            return

        scanners, err = discover_scanners_naps2()
        if err:
            self.Log(err, engine_event=True)
        if not scanners:
            msg = "No NAPS2 scanner devices were detected."
            if err:
                msg += f"\n\nDetails: {err}"
            wx.MessageBox(msg, "Scan via NAPS2", wx.OK | wx.ICON_INFORMATION)
            return

        dlg = ScanSettingsDialog(self, scanners, self.cfg, script_dir=SCRIPT_DIR)
        if dlg.ShowModal() != wx.ID_OK:
            dlg.Destroy()
            return
        try:
            settings = dlg.GetSettings()
        except ValueError as ex:
            dlg.Destroy()
            wx.MessageBox(str(ex), "Scan via NAPS2", wx.OK | wx.ICON_ERROR)
            return
        finally:
            dlg.Destroy()

        scanner = settings["scanner"]
        dpi = settings["dpi"]
        output_dir = settings["output_dir"]
        preset_label = settings["preset_label"]
        self.cfg = shared_apply_scan_settings(
            self.cfg,
            scanner=scanner,
            dpi=dpi,
            output_dir=output_dir,
            preset_label=preset_label,
            driver_from_source_fn=_driver_from_scanner_source,
        )
        save_json(CONFIG_FILE, self.cfg)
        self.SaveActiveSession()

        self.SetRunningState(True)
        self.Log(shared_build_scan_start_message(scanner, dpi, preset_label), engine_event=True)
        threading.Thread(
            target=self.ScanWorkerNaps2,
            args=(dict(scanner), dpi, output_dir),
            daemon=True
        ).start()

    def _attempt_naps2_scan(self, base_cmd, scanner_name, driver, dpi, output_pdf):
        base_opts = ["--output", output_pdf, "--dpi", str(dpi)]
        if driver:
            base_opts.extend(["--driver", driver])
        if scanner_name:
            base_opts.extend(["--device", scanner_name])

        attempts = [
            [*base_cmd, "scan", *base_opts],
            [*base_cmd, "scan", "-o", output_pdf, *base_opts[2:]],
            [*base_cmd, *base_opts],
            [*base_cmd, "-o", output_pdf, *base_opts[2:]],
        ]
        tried = set()
        errors = []
        for cmd in attempts:
            key = tuple(cmd)
            if key in tried:
                continue
            tried.add(key)
            code, out, err = _run_command_capture(cmd, timeout=300)
            if code == 0:
                return True, cmd, out, err
            detail = err.strip() or out.strip() or f"exit {code}"
            errors.append(f"{' '.join(cmd)} -> {detail}")
        return False, None, "", "\n".join(errors[:3])

    def ScanWorkerNaps2(self, scanner, dpi, output_dir):
        try:
            os.makedirs(output_dir, exist_ok=True)
            scan_session = shared_begin_scan_session(output_dir)
            scanner_name = scanner.get("name", "").strip()
            driver = shared_resolve_scan_driver(
                scanner,
                self.cfg,
                driver_from_source_fn=_driver_from_scanner_source,
            )

            commands = shared_choose_scan_commands(_candidate_naps2_commands())
            scan_result = shared_execute_scan_commands(
                commands,
                attempt_scan_fn=self._attempt_naps2_scan,
                scanner_name=scanner_name,
                driver=driver,
                dpi=dpi,
                output_pdf=scan_session["output_pdf"],
                log_cb=lambda message: self.Log(message, engine_event=True),
            )
            if not scan_result["success"]:
                if scan_result["error_kind"] == "missing_executable":
                    self.Log("NAPS2 scan failed: executable not found.", engine_event=True)
                    wx.CallAfter(
                        wx.MessageBox,
                        scan_result["details"],
                        "Scan via NAPS2",
                        wx.OK | wx.ICON_ERROR,
                    )
                    return
                self.Log(f"NAPS2 scan command failed. {scan_result['details']}", engine_event=True)
                wx.CallAfter(
                    wx.MessageBox,
                    "NAPS2 did not complete the scan command.\n\n"
                    "Check scanner power/cable/driver and try again.\n\n"
                    f"Details: {scan_result['details'][:500]}",
                    "Scan via NAPS2",
                    wx.OK | wx.ICON_ERROR,
                )
                return

            new_files = shared_resolve_scan_output_files(
                output_dir,
                scan_session["before_paths"],
                scan_session["started_ts"],
                scan_session["output_pdf"],
                collect_scan_files_fn=_collect_scan_files,
            )
            if not new_files:
                self.Log("NAPS2 completed but no new scan files were found.", engine_event=True)
                wx.CallAfter(
                    wx.MessageBox,
                    "Scan command completed, but no new output file was detected.",
                    "Scan via NAPS2",
                    wx.OK | wx.ICON_WARNING,
                )
                return

            wx.CallAfter(
                self._ImportScannedFilesToQueue,
                list(new_files),
                output_dir,
            )
        finally:
            wx.CallAfter(self.SetRunningState, False)
            wx.CallAfter(self.file_list.SetFocus)

    def _AutoStartAfterScan(self, force_merge_output=False):
        if self.is_running:
            return
        if not self.queue:
            return
        self.pending_scan_merge_extract = bool(force_merge_output)
        self.Log("Auto-starting extraction for queued files after scan.", engine_event=True)
        self.OnStart(None)

    def _ImportScannedFilesToQueue(self, paths, output_dir):
        added_paths = self.AddPathEntries(paths, source_root=output_dir)
        self.RefreshQueue()
        self.SelectQueueRowsByPaths(added_paths)
        self.SaveActiveSession()
        dialog_message, log_message = shared_build_scan_completion_message(paths, added_paths)
        self.Log(log_message, engine_event=True)
        wx.MessageBox(
            dialog_message,
            "Scan via NAPS2",
            wx.OK | wx.ICON_INFORMATION,
        )

    def OnMacAccessCheck(self, e):
        self.PromptMacFolderAccess(force=True)

    def OnOpenHelp(self, e):
        help_path = os.path.abspath(HELP_FILE)
        if not os.path.exists(help_path):
            wx.MessageBox(
                f"Help file not found:\n{help_path}\n\nBuild: {self.build_stamp}",
                "Help File Missing",
                wx.OK | wx.ICON_WARNING,
            )
            return
        try:
            wx.LaunchDefaultBrowser(f"file://{help_path}")
            self.Log(f"Opened help guide: {help_path} (Build {self.build_stamp})")
        except Exception:
            try:
                webbrowser.open(f"file://{help_path}", new=2)
                self.Log(f"Opened help guide: {help_path} (Build {self.build_stamp})")
            except Exception as ex:
                wx.MessageBox(
                    f"Unable to open help file.\n\nDetails: {ex}",
                    "Help Launch Error",
                    wx.OK | wx.ICON_ERROR,
                )

    def OnAboutBuild(self, e):
        self.RefreshLicenseStatus()
        wx.MessageBox(
            (
                f"{APP_NAME}\n\n"
                f"Build: {self.build_stamp}\n"
                f"Runtime: {sys.executable if getattr(sys, 'frozen', False) else __file__}\n"
                f"Python: {sys.version.split()[0]}\n"
                f"wx: {wx.version()}\n\n"
                f"{self.license_status_text}"
            ),
            "About Build",
            wx.OK | wx.ICON_INFORMATION,
        )

    def OpenExternalLink(self, url, label):
        try:
            wx.LaunchDefaultBrowser(url)
            self.Log(f"Opened {label}: {url}")
        except Exception:
            try:
                webbrowser.open(url, new=2)
                self.Log(f"Opened {label}: {url}")
            except Exception as ex:
                wx.MessageBox(
                    f"Unable to open link.\n\nDetails: {ex}",
                    "Link Launch Error",
                    wx.OK | wx.ICON_ERROR,
                )

    def OnDonate(self, e):
        dlg = wx.SingleChoiceDialog(
            self,
            "If you'd like to support the developer, donations help cover ongoing testing and API usage costs.\n\nChoose a donation method:",
            "Support Chronicle",
            ["Buy Me a Coffee", "PayPal"],
        )
        if dlg.ShowModal() == wx.ID_OK:
            choice = dlg.GetStringSelection()
            if choice == "Buy Me a Coffee":
                self.OpenExternalLink(DONATE_BUYMEACOFFEE_URL, "Buy Me a Coffee")
            elif choice == "PayPal":
                self.OpenExternalLink(DONATE_PAYPAL_URL, "PayPal")
        dlg.Destroy()

    def OnPrefs(self, e):
        self.RefreshLicenseStatus()
        d = PrefsDialog(
            self,
            self.cfg,
            config_file=CONFIG_FILE,
            save_json=save_json,
            get_translation_target=get_translation_target,
            profile_tooltip_text=profile_tooltip_text,
            apply_profile_preset=apply_profile_preset,
            license_status_text=self.license_status_text,
            import_license_handler=self.OnImportLicense,
        )
        d.ShowModal()
        d.Destroy()
        self.cfg = load_json(CONFIG_FILE, self.cfg)
        if not self.is_running:
            self.cfg["model_name"] = self.ResolveModelNameForAvailableKeys(
                self.GetPreferredModelName(self.cfg.get("doc_profile", "standard"))
            )
            for row in self.queue:
                self.NormalizeRowSettings(row)
            self.RefreshQueue()
        self._sync_controls_from_cfg()
        self.SaveActiveSession()

    def RefreshLicenseStatus(self):
        self.license_validation = None
        self.license_public_key = None
        self.license_public_key_error = ""

        if (
            shared_format_license_status is None
            or shared_install_license_file is None
            or shared_load_installed_license is None
            or shared_resolve_public_key is None
        ):
            self.license_status_text = (
                "License system unavailable\n\n"
                "The optional licensing dependency is not installed in this Python environment."
            )
            return self.license_status_text

        try:
            self.license_public_key = shared_resolve_public_key(
                app_data_dir=APP_DATA_DIR,
                script_dir=SCRIPT_DIR,
            )
        except Exception as ex:
            self.license_public_key_error = f"Unable to load the Chronicle public verification key.\n\nDetails: {ex}"

        if self.license_public_key_error:
            self.license_status_text = shared_format_license_status(
                None,
                public_key_available=False,
                public_key_error=self.license_public_key_error,
            )
            return self.license_status_text

        if self.license_public_key is not None:
            try:
                self.license_validation = shared_load_installed_license(
                    app_data_dir=APP_DATA_DIR,
                    public_key=self.license_public_key,
                )
            except Exception as ex:
                self.license_public_key_error = f"Unable to read the installed Chronicle license.\n\nDetails: {ex}"

        self.license_status_text = shared_format_license_status(
            self.license_validation,
            public_key_available=self.license_public_key is not None,
            public_key_error=self.license_public_key_error,
        )
        return self.license_status_text

    def OnImportLicense(self, e=None):
        self.RefreshLicenseStatus()
        if shared_install_license_file is None or shared_resolve_public_key is None:
            wx.MessageBox(
                self.license_status_text,
                "License Import Unavailable",
                wx.OK | wx.ICON_WARNING,
            )
            return self.license_status_text

        if self.license_public_key is None:
            wx.MessageBox(
                self.license_status_text,
                "License Verification Key Missing",
                wx.OK | wx.ICON_WARNING,
            )
            return self.license_status_text

        dlg = wx.FileDialog(
            self,
            "Import Chronicle License",
            wildcard="Chronicle license (*.json)|*.json|All files (*.*)|*.*",
            style=wx.FD_OPEN | wx.FD_FILE_MUST_EXIST,
        )
        try:
            if dlg.ShowModal() != wx.ID_OK:
                return self.license_status_text

            selected_path = dlg.GetPath()
        finally:
            dlg.Destroy()

        try:
            result = shared_install_license_file(
                selected_path,
                app_data_dir=APP_DATA_DIR,
                public_key=self.license_public_key,
            )
        except Exception as ex:
            wx.MessageBox(
                f"Unable to import the selected license file.\n\nDetails: {ex}",
                "License Import Error",
                wx.OK | wx.ICON_ERROR,
            )
            return self.RefreshLicenseStatus()

        self.RefreshLicenseStatus()
        if result.valid and result.license_data:
            licensed_to = result.license_data.get("issued_to", "Unknown")
            tier = result.license_data.get("tier", "unknown")
            self.Log(f"Installed Chronicle license for {licensed_to} ({tier}).")
            wx.MessageBox(
                self.license_status_text,
                "License Imported",
                wx.OK | wx.ICON_INFORMATION,
            )
        else:
            wx.MessageBox(
                self.license_status_text,
                "License Not Accepted",
                wx.OK | wx.ICON_WARNING,
            )
        return self.license_status_text

    def GetClient(self, model):
        vendor = "anthropic" if "claude" in model else "openai" if "gpt" in model else "gemini"
        key_vendor = "claude" if vendor == "anthropic" else vendor
        key = self.keys.get(key_vendor)
        cache_key = f"{vendor}:{key}"
        if cache_key in self.clients:
            return self.clients[cache_key]
        if vendor == "anthropic":
            import anthropic
            client = anthropic.Anthropic(api_key=key)
        elif vendor == "openai":
            import openai
            client = openai.OpenAI(api_key=key)
        else:
            client = genai.Client(api_key=key)
        self.clients[cache_key] = client
        return client

    def OnStart(self, e):
        if self.is_running:
            return
        if not self.queue:
            wx.MessageBox("File list is empty. Add files first.", "No Files", wx.OK | wx.ICON_INFORMATION)
            return

        start_context = shared_begin_run_start(
            resume_incomplete_only=self.resume_incomplete_only,
            scheduled_start_ts=self.scheduled_start_ts,
        )
        resume_mode = start_context["resume_mode"]
        self.resume_incomplete_only = start_context["next_resume_incomplete_only"]
        if start_context["should_clear_schedule"]:
            self._set_scheduled_start(None)
            self.Log("Cleared scheduled extraction because extraction started now.", engine_event=True)

        if resume_mode:
            queue_state = shared_prepare_queue_for_start(
                self.queue,
                resume_mode=True,
                normalize_row_settings_fn=self.NormalizeRowSettings,
            )
            self.is_paused = queue_state.get("is_paused", self.is_paused)
            self._sync_controls_from_cfg()
        else:
            control_settings = self.BuildSettingsFromControls()
            force_merge_from_scan = bool(self.pending_scan_merge_extract)
            self.pending_scan_merge_extract = False
            self.cfg = shared_apply_start_configuration(
                self.cfg,
                control_settings=control_settings,
                force_merge_from_scan=force_merge_from_scan,
                recursive_scan=self.choice_recursive.GetSelection() == 1,
                dest_mode=self.dest_choice.GetSelection(),
                custom_dest=self.txt_dest.GetValue(),
                preserve_source_structure=self.chk_preserve_structure.GetValue(),
                delete_source_on_success=self.chk_delete_originals.GetValue(),
                script_dir=SCRIPT_DIR,
            )
            if force_merge_from_scan:
                self.Log("Scan setting applied: extraction merge mode forced ON for this run.", engine_event=True)
            output_error = shared_validate_output_destination(self.cfg)
            if output_error:
                title = "Missing Output Folder" if "Please choose" in output_error else "Invalid Output Folder"
                wx.MessageBox(output_error, title, wx.OK | wx.ICON_ERROR)
                return
            save_json(CONFIG_FILE, self.cfg)
            shared_prepare_queue_for_start(
                self.queue,
                resume_mode=False,
                normalize_row_settings_fn=self.NormalizeRowSettings,
            )
            split_result = shared_expand_multi_range_pdf_rows(
                self.queue,
                normalize_row_settings_fn=self.NormalizeRowSettings,
                normalize_pdf_page_scope_text_fn=normalize_pdf_page_scope_text,
            )
            if split_result.get("changed"):
                self.RefreshQueue()
                self.SaveActiveSession()
                self.Log(
                    "Expanded multi-range PDF page selections into separate queued outputs.",
                    engine_event=True,
                )

        pending_rows = shared_collect_pending_rows(self.queue)
        if not pending_rows:
            wx.MessageBox("No queued tasks to process. Resume paused tasks or add files.", "Nothing Queued", wx.OK | wx.ICON_INFORMATION)
            self.RefreshQueue()
            self.SaveActiveSession()
            return

        pdf_scope_error = shared_validate_pending_pdf_page_scopes(
            pending_rows,
            normalize_row_settings_fn=self.NormalizeRowSettings,
            pdf_reader_factory=PdfReader,
            normalize_pdf_page_scope_text_fn=normalize_pdf_page_scope_text,
            parse_pdf_page_scope_spec_fn=parse_pdf_page_scope_spec,
        )
        if pdf_scope_error:
            wx.MessageBox(
                f"PDF page scope is invalid for {pdf_scope_error['filename']}.\n\nDetails: {pdf_scope_error['details']}",
                "Invalid PDF Pages",
                wx.OK | wx.ICON_ERROR,
            )
            return

        if self.cfg.get('delete_source_on_success', False):
            source_dirs = sorted({os.path.dirname(row.get('path', '')) for row in pending_rows if row.get('path')})
            folder_count = len(source_dirs)
            folder_phrase = f"{folder_count} folder(s)" if folder_count else "the queued source folders"
            output_plan = (
                "Results are set to save beside each source file."
                if int(self.cfg.get('dest_mode', 0)) == 0
                else (
                    "Results are set to save in the custom output folder while preserving the scanned folder structure."
                    if bool(self.cfg.get('preserve_source_structure', True))
                    else "Results are set to save in one custom output folder without preserving the scanned folder structure."
                )
            )
            recursive_notice = " Add Folder is currently set to include nested subfolders." if bool(self.cfg.get('recursive_scan', False)) else ""
            confirm = wx.MessageBox(
                "Delete Originals is enabled.\n\n"
                f"Chronicle will permanently delete each source file after successful extraction across {folder_phrase}.{recursive_notice}\n\n"
                f"{output_plan}\n\n"
                "Do you want to continue?",
                "Confirm Source Deletion",
                wx.YES_NO | wx.NO_DEFAULT | wx.ICON_WARNING,
            )
            if confirm != wx.YES:
                self.Log("Extraction cancelled by user (delete originals confirmation).", engine_event=True)
                return

        missing_key = shared_find_missing_api_key_requirement(
            pending_rows,
            normalize_row_settings_fn=self.NormalizeRowSettings,
            model_from_label_fn=self.ModelFromLabel,
            label_from_model_fn=self.LabelFromModel,
            has_vendor_key_fn=self.HasProviderKey,
        )
        if missing_key:
            wx.MessageBox(
                f"Missing API Key for {missing_key['vendor'].upper()} (required by {missing_key['label']}).",
                "Error",
                wx.OK | wx.ICON_ERROR,
            )
            return

        self.RefreshQueue()
        self.RefreshQueueWorkUnitEstimates()
        self.SaveActiveSession()
        try:
            selected_profile_label = PROFILE_KEY_TO_LABEL.get(
                self.cfg.get("doc_profile", "standard"),
                "Standard",
            )
            shared_update_continuity_runtime_status(
                script_dir=SCRIPT_DIR,
                event="run_start",
                detail=(
                    f"Queued {len(pending_rows)} file(s); "
                    f"profile {selected_profile_label}; "
                    f"format {self.cfg.get('format_type', 'html').upper()}."
                ),
            )
        except Exception as ex:
            self.Log(f"Warning: could not update continuity runtime status ({ex})", engine_event=True)
        self.lt.Clear()
        for key, value in shared_build_run_reset_state().items():
            setattr(self, key, value)
        self.current_run_resume_mode = bool(resume_mode)
        self.current_file_name = ""
        self.current_run_started_ts = time.time()
        self._last_engine_activity_ts = self.current_run_started_ts
        self._last_progress_heartbeat_ts = 0.0
        self.SetRunningState(True)
        start_messages = shared_build_start_messages(resume_mode=resume_mode, pending_count=len(pending_rows))
        self.Log(start_messages["log_message"], engine_event=True)
        self.SafeSetStatusText(start_messages["status_text"])
        threading.Thread(target=self.Worker, daemon=True).start()

    def OnScheduleExtraction(self, e):
        if self.is_running:
            return
        if not self.queue:
            wx.MessageBox("Queue at least one file before scheduling extraction.", "Schedule Extraction", wx.OK | wx.ICON_INFORMATION)
            return
        dlg = ScheduleExtractionDialog(self, existing_ts=self.scheduled_start_ts)
        try:
            result = dlg.ShowModal()
            if result == ScheduleExtractionDialog.CLEAR_SCHEDULE:
                if self.scheduled_start_ts is not None:
                    self._set_scheduled_start(None)
                    self.Log("Cleared scheduled extraction.", engine_event=True)
                return
            if result != wx.ID_OK:
                return
            try:
                scheduled_ts = dlg.GetScheduledTimestamp()
            except ValueError as ex:
                wx.MessageBox(str(ex), "Schedule Extraction", wx.OK | wx.ICON_ERROR)
                return
            if scheduled_ts <= time.time():
                wx.MessageBox("Scheduled time must be in the future.", "Schedule Extraction", wx.OK | wx.ICON_ERROR)
                return
            self._set_scheduled_start(scheduled_ts)
            self.Log(
                f"Scheduled extraction for {self._format_timestamp_local(scheduled_ts)} (local time).",
                engine_event=True,
            )
        finally:
            dlg.Destroy()

    def Worker(self):
        try:
            resume_mode = bool(getattr(self, "current_run_resume_mode", False))
            jobs = self.GetOrderedJobsForProcessing()
            processing_log = lambda message: self.Log(message, engine_event=True)
            run_plan = shared_build_worker_run_plan(
                self.cfg,
                jobs,
                normalize_row_settings_fn=self.NormalizeRowSettings,
                streamable_formats=STREAMABLE_FORMATS,
            )
            custom_dest = run_plan["custom_dest"]
            dest_mode = run_plan["dest_mode"]
            merge_mode = run_plan["merge_mode"]
            low_memory_mode = run_plan["low_memory_mode"]
            memory_telemetry = run_plan["memory_telemetry"]
            queued_jobs = run_plan["queued_jobs"]
            queued_formats = run_plan["queued_formats"]
            default_fmt = run_plan["default_fmt"]
            merge_fmt = run_plan["merge_fmt"]
            streamable_fmt = run_plan["streamable_fmt"]
            master_output_path = None
            master_temp_path = None
            master_progress_temp_path = None
            master_resume_state_path = None
            master_progress_file_obj = None
            master_file_obj = None
            master_memory = run_plan["master_memory"]
            completed_merged_job_paths = set()
            for message in run_plan["messages"]:
                self.Log(message, engine_event=True)

            if memory_telemetry:
                self.Log(f"[Memory] Peak RSS at run start: {get_peak_rss_mb():.1f} MB", engine_event=True)

            if dest_mode == 1:
                os.makedirs(custom_dest, exist_ok=True)
            if merge_mode:
                first_job_cfg = dict(self.cfg)
                first_job_settings = self.NormalizeRowSettings(queued_jobs[0] if queued_jobs else {})
                first_job_cfg.update(first_job_settings)
                master_output_path = self.ResolveMergeOutputPath(jobs, merge_fmt, custom_dest, dest_mode)
                master_progress_temp_path = master_output_path + ".progress.txt.tmp"
                if streamable_fmt:
                    master_temp_path = master_output_path + ".tmp"
                merge_resume = {"recovered_units": 0, "original_total_units": len(queued_jobs), "completed_job_paths": []}
                if resume_mode:
                    merge_resume = shared_load_merge_resume_state(
                        progress_temp_path=master_progress_temp_path,
                        expected_total_units=len(queued_jobs),
                        path_exists_fn=os.path.exists,
                        open_fn=open,
                    )
                master_resume_state_path = merge_resume.get("resume_state_path")
                completed_merged_job_paths = set(merge_resume.get("completed_job_paths", []))
                preserve_merge_progress = bool(
                    resume_mode
                    and merge_resume.get("recovered_units", 0) > 0
                    and os.path.exists(master_progress_temp_path)
                    and (not streamable_fmt or os.path.exists(master_temp_path))
                )
                if os.path.exists(master_progress_temp_path) and not preserve_merge_progress:
                    os.remove(master_progress_temp_path)
                if preserve_merge_progress:
                    ensure_progress_sidecar_header(master_progress_temp_path)
                master_progress_file_obj = open(master_progress_temp_path, "a" if preserve_merge_progress else "w", encoding="utf-8")
                if not preserve_merge_progress:
                    master_progress_file_obj.write(build_progress_state_header({}))
                    master_progress_file_obj.flush()
                self.Log(f"[Progress] In-progress temp output: {master_progress_temp_path}", engine_event=True)
                if streamable_fmt:
                    if os.path.exists(master_temp_path) and not preserve_merge_progress:
                        os.remove(master_temp_path)
                    master_file_obj = MirroredTextWriter(
                        open(master_temp_path, "a" if preserve_merge_progress else "w", encoding="utf-8"),
                        master_progress_file_obj,
                    )
                    if not preserve_merge_progress:
                        write_header(
                            master_file_obj,
                            "Chronicle Merged",
                            merge_fmt,
                            get_output_lang_code(first_job_cfg),
                            get_output_text_direction(first_job_cfg),
                        )
                else:
                    master_memory = MirroredProgressMemory(
                        progress_file_obj=master_progress_file_obj,
                        progress_temp_path=master_progress_temp_path,
                        retain_chunks=False,
                    )
                if preserve_merge_progress:
                    total_merge_units = int(merge_resume.get("original_total_units") or len(queued_jobs) or 0)
                    recovered_merge_units = int(merge_resume.get("recovered_units", 0) or 0)
                    self.Log(
                        f"[Resume] Found preserved merged progress. Recovered {recovered_merge_units} of "
                        f"{total_merge_units or '?'} files; continuing with the remaining files only.",
                        engine_event=True,
                    )
                    for queued_job in queued_jobs:
                        if queued_job.get("path") in completed_merged_job_paths:
                            self.SetQueueStatus(queued_job.get("_queue_index", 0), "Done")
                self.Log(f"[Merge] Seamless merge enabled -> {master_output_path}", engine_event=True)

            active_file_counter = 0
            for job in jobs:
                if self.stop_requested:
                    self.Log("Stop requested. Ending extraction run safely.", engine_event=True)
                    break
                self.WaitWhilePaused()
                qidx = job.get('_queue_index', 0)
                self.current_processing_index = qidx
                current_status = str(self.queue[qidx].get("status", "Queued")) if 0 <= qidx < len(self.queue) else "Queued"
                if merge_mode and job.get("path") in completed_merged_job_paths:
                    self.SetQueueStatus(qidx, "Done")
                    continue
                if current_status in SESSION_TERMINAL_STATUSES:
                    continue
                if current_status == "Paused":
                    self.Log(f"Deferred paused task: {os.path.basename(job['path'])}", engine_event=True)
                    continue
                prepared = shared_prepare_job_execution_context(
                    job,
                    cfg=self.cfg,
                    resume_mode=resume_mode,
                    low_memory_mode=low_memory_mode,
                    low_memory_pdf_audit_skip_mb=LOW_MEMORY_PDF_AUDIT_SKIP_MB,
                    custom_dest=custom_dest,
                    dest_mode=dest_mode,
                    merge_mode=merge_mode,
                    master_output_path=master_output_path,
                    master_temp_path=master_temp_path,
                    master_file_obj=master_file_obj,
                    master_memory=master_memory,
                    streamable_formats=STREAMABLE_FORMATS,
                    supported_extensions=SUPPORTED_EXTENSIONS,
                    normalize_row_settings_fn=self.NormalizeRowSettings,
                    build_prompt_fn=build_prompt,
                    model_from_label_fn=self.ModelFromLabel,
                    get_client_fn=self.GetClient,
                    determine_needs_pdf_audit_fn=shared_determine_needs_pdf_audit,
                    compute_target_dir_fn=shared_compute_target_dir,
                    resolve_output_path_fn=shared_resolve_output_path,
                    write_header_fn=write_header,
                    get_output_lang_code_fn=get_output_lang_code,
                    get_output_text_direction_fn=get_output_text_direction,
                    pdf_reader_factory=PdfReader,
                    normalize_pdf_page_scope_text_fn=normalize_pdf_page_scope_text,
                    parse_pdf_page_scope_spec_fn=parse_pdf_page_scope_spec,
                    set_queue_status_fn=self.SetQueueStatus,
                    log_cb=processing_log,
                )
                if prepared['skip']:
                    continue

                fp = prepared['path']
                fn = prepared['file_name']
                base = prepared['base']
                ext = prepared['ext']
                job_cfg = prepared['job_cfg']
                fmt = prepared['fmt']
                prompt = prepared['prompt']
                model = prepared['model']
                routing_mode = prepared.get('routing_mode', 'automatic')
                routing_reason = prepared.get('routing_reason', '')
                client = prepared['client']
                needs_pdf_audit = prepared['needs_pdf_audit']
                output_path = prepared['output_path']
                temp_path = prepared['temp_path']
                progress_temp_path = prepared.get('progress_temp_path')
                resume_state_path = prepared.get('resume_state_path')
                progress_file_obj = prepared.get('progress_file_obj')
                file_obj = prepared['file_obj']
                memory = prepared['memory']
                recovered_units = int(prepared.get('recovered_units', 0) or 0)
                original_total_units = int(prepared.get('original_total_units') or 0)
                resume_from_unit = int(prepared.get('resume_from_unit', 0) or 0)

                self.SetQueueStatus(qidx, "Processing")
                active_file_counter += 1
                self.current_file_ordinal = active_file_counter
                self.current_file_name = fn
                self.current_file_resume_state_path = resume_state_path
                self.current_file_page_done = recovered_units
                self.current_file_resume_recovered_units = recovered_units
                unit_estimate = shared_estimate_current_file_total_units(
                    ext,
                    fp,
                    job_cfg,
                    pdf_reader_factory=PdfReader,
                    normalize_pdf_page_scope_text_fn=normalize_pdf_page_scope_text,
                    parse_pdf_page_scope_spec_fn=parse_pdf_page_scope_spec,
                    pptx_slide_count_fn=lambda pptx_path: len(__import__('pptx').Presentation(pptx_path).slides),
                    estimate_text_work_units_fn=lambda path, ext, cfg: shared_estimate_text_work_units(
                        path,
                        ext,
                        text_chunk_chars=TEXT_CHUNK_CHARS,
                        csv_to_accessible_text_fn=csv_to_accessible_text,
                        clean_text_fn=clean_text,
                        batch_text_chunks_fn=batch_text_chunks,
                        docx_module=docx,
                        openpyxl_module=openpyxl,
                    ),
                )
                self.current_file_page_total = original_total_units or unit_estimate["total_units"]
                self.current_file_resume_remaining_units = max(0, self.current_file_page_total - recovered_units)
                self.current_file_unit_label = unit_estimate.get("unit_label", "items")
                if recovered_units > 0 and self.current_file_resume_remaining_units > 0:
                    self.Log(
                        f"[Resume] Recovered previous progress for {fn}: {recovered_units}/{self.current_file_page_total} "
                        f"{self.current_file_unit_label} already complete. Resuming the remaining "
                        f"{self.current_file_resume_remaining_units} {self.current_file_unit_label}.",
                        engine_event=True,
                    )
                if unit_estimate["selected_scope"]:
                    self.Log(
                        f"[Scope] {fn}: limiting PDF reading to pages {unit_estimate['selected_scope']} (selected {unit_estimate['selected_count']} of {unit_estimate['source_total']}).",
                        engine_event=True,
                    )
                self.Log(
                    f"Processing file {self.current_file_ordinal}/{len(queued_jobs)}: {fn} with {self.LabelFromModel(model)} "
                    f"({fmt.upper()}, {PROFILE_KEY_TO_LABEL.get(job_cfg.get('doc_profile', 'standard'), 'Standard')}). "
                    f"{self.current_file_page_done}/{self.current_file_page_total} {self.current_file_unit_label} processed.",
                    engine_event=True,
                )
                if routing_mode == "automatic" and routing_reason:
                    self.Log(f"[Auto Engine] {routing_reason}", engine_event=True)
                if memory_telemetry:
                    self.Log(f"[Memory] Peak RSS before task: {get_peak_rss_mb():.1f} MB", engine_event=True)
                try:
                    shared_process_job_content(
                        ext,
                        cfg=self.cfg,
                        path=fp,
                        file_name=fn,
                        temp_path=temp_path,
                        fmt=fmt,
                        prompt=prompt,
                        model=model,
                        client=client,
                        file_obj=file_obj,
                        memory=memory,
                        processing_log=processing_log,
                        pause_cb=self.WaitWhilePaused,
                        page_scope=job_cfg.get('pdf_page_scope', ''),
                        describe_quality_score_fn=describe_quality_score,
                        assess_image_file_quality_fn=assess_image_file_quality,
                        update_progress_state_fn=lambda done_pages, total_pages, recovered=recovered_units: self._update_pdf_progress_state(
                            done_pages + recovered,
                            total_pages + recovered,
                        ),
                        should_log_page_progress_fn=self.ShouldLogPageProgress,
                        refresh_progress_fn=lambda: wx.CallAfter(self.UpdateProgressIndicators),
                        persist_progress_state_fn=self._persist_resume_state if ext in ['.pdf', '.docx', '.txt', '.md', '.rtf', '.csv', '.js', '.xlsx', '.xls', '.epub', '.pptx', '.ppt'] else None,
                        resume_from_unit=resume_from_unit,
                        needs_pdf_audit=needs_pdf_audit,
                        append_pdf_audit_appendix_if_needed_fn=shared_append_pdf_audit_appendix_if_needed,
                        run_pdf_textlayer_audit_fn=run_pdf_textlayer_audit,
                        render_audit_appendix_fn=render_audit_appendix,
                        append_generated_text_fn=append_generated_text,
                        coverage_warn_threshold=PDF_TEXTLAYER_AUDIT_COVERAGE_WARN,
                        coverage_append_full_threshold=PDF_TEXTLAYER_AUDIT_COVERAGE_APPEND_FULL,
                        process_pdf_fn=process_pdf,
                        process_pptx_fn=process_pptx,
                        process_epub_fn=process_epub,
                        process_img_fn=process_img,
                        process_text_fn=process_text,
                        original_total_units=self.current_file_page_total,
                    )

                    success = shared_finalize_job_success(
                        merge_mode=merge_mode,
                        job_cfg=job_cfg,
                        temp_path=temp_path,
                        memory=memory,
                        base=base,
                        file_obj=file_obj,
                        fmt=fmt,
                        output_path=output_path,
                        source_path=fp,
                        file_name=fn,
                        ext=ext,
                        current_file_page_total=self.current_file_page_total,
                        memory_telemetry=memory_telemetry,
                        delete_source_on_success=bool(self.cfg.get('delete_source_on_success', False)),
                        dispatch_save_fn=dispatch_save,
                        write_footer_fn=write_footer,
                        cleanup_output_text_fn=cleanup_output_text,
                        is_protected_path_fn=is_path_within_protected_input_dirs,
                        set_queue_status_fn=lambda status: self.SetQueueStatus(qidx, status),
                        log_cb=processing_log,
                        get_peak_rss_mb_fn=get_peak_rss_mb,
                        progress_temp_path=progress_temp_path,
                        progress_file_obj=progress_file_obj,
                        resume_state_path=resume_state_path,
                    )
                    if success['set_page_done_to_total']:
                        self.current_file_page_done = self.current_file_page_total
                        self.total_pages_processed += success['page_total_increment']
                    if merge_mode and master_progress_temp_path:
                        completed_merged_job_paths.add(fp)
                        self._persist_merge_resume_state(
                            master_progress_temp_path,
                            sorted(completed_merged_job_paths),
                            len(queued_jobs),
                        )
                except StopRequestedError:
                    self.SetQueueStatus(qidx, "Stopped")
                    self.Log(f"Stopped active task: {fn}", engine_event=True)
                    break
                except Exception as ex:
                    shared_handle_job_error(
                        merge_mode=merge_mode,
                        file_obj=file_obj,
                        temp_path=temp_path,
                        file_name=fn,
                        error=ex,
                        set_queue_status_fn=lambda status: self.SetQueueStatus(qidx, status),
                        log_cb=processing_log,
                        progress_temp_path=progress_temp_path,
                        progress_file_obj=progress_file_obj,
                        resume_state_path=resume_state_path,
                    )

            if merge_mode and master_output_path:
                shared_finalize_merged_output(
                    cfg=self.cfg,
                    merge_fmt=merge_fmt,
                    streamable_fmt=streamable_fmt,
                    master_file_obj=master_file_obj,
                    master_temp_path=master_temp_path,
                    master_output_path=master_output_path,
                    master_memory=master_memory,
                    write_footer_fn=write_footer,
                    cleanup_output_text_fn=cleanup_output_text,
                    strip_synthetic_headings_fn=strip_synthetic_page_filename_headings,
                    dispatch_save_fn=dispatch_save,
                    log_cb=processing_log,
                    progress_temp_path=master_progress_temp_path,
                    progress_file_obj=master_progress_file_obj,
                    resume_state_path=master_resume_state_path,
                )

            shared_finalize_worker_completion(
                auto_save_processing_log_fn=self.AutoSaveProcessingLog,
                log_cb=processing_log,
                platform_system=platform.system(),
                subprocess_popen=subprocess.Popen,
                winsound_module=winsound,
            )
        except Exception as ex:
            self.Log(f"FATAL: {ex}", engine_event=True)
        finally:
            try:
                status_counts = {}
                for row in self.queue:
                    label = str(row.get("status", "Queued"))
                    status_counts[label] = status_counts.get(label, 0) + 1
                ordered_labels = [
                    "Done",
                    "Error",
                    "Paused",
                    "Stopped",
                    "Queued",
                    "Processing",
                    "Skipped",
                    "Missing",
                    "Unsupported",
                ]
                parts = []
                for label in ordered_labels:
                    count = status_counts.pop(label, 0)
                    if count:
                        parts.append(f"{label} {count}")
                for label in sorted(status_counts):
                    parts.append(f"{label} {status_counts[label]}")
                shared_update_continuity_runtime_status(
                    script_dir=SCRIPT_DIR,
                    event="run_complete",
                    detail="Run finished. " + (", ".join(parts) if parts else "No queue items recorded."),
                )
            except Exception as ex:
                self.Log(f"Warning: could not finalize continuity runtime status ({ex})", engine_event=True)
            shared_finalize_worker_session(
                master_file_obj=locals().get('master_file_obj'),
                has_incomplete_items=self._has_incomplete_items(self.queue),
                save_active_session_fn=self.SaveActiveSession,
                delete_active_session_fn=self.DeleteActiveSession,
                set_running_state_fn=lambda value: wx.CallAfter(self.SetRunningState, value),
            )
            if self.autorun_active and self.autorun_close_on_finish:
                wx.CallAfter(self.Close)

def _termination_guard(signum, frame):
    global ACTIVE_FRAME
    try:
        if ACTIVE_FRAME and ACTIVE_FRAME.is_running:
            ACTIVE_FRAME.is_paused = True
            shared_pause_current_processing_row(ACTIVE_FRAME.queue, ACTIVE_FRAME.current_processing_index)
            ACTIVE_FRAME.SaveActiveSession()
            print(f"[Guard] Ignored termination signal {signum} while extraction is active. Session checkpoint saved.")
            return
    except Exception as ex:
        print(f"[Guard] Termination guard warning: {ex}")
    raise SystemExit(0)

if __name__ == '__main__':
    if hasattr(signal, "SIGTERM"):
        signal.signal(signal.SIGTERM, _termination_guard)
    if hasattr(signal, "SIGINT"):
        signal.signal(signal.SIGINT, _termination_guard)
    shared_emit_launch_continuity(script_dir=os.path.dirname(os.path.abspath(__file__)))
    app = wx.App(False)
    frame = MainFrame()
    frame.Show()
    app.MainLoop()
